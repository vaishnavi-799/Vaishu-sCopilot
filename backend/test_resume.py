import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from fastapi.testclient import TestClient
from bson import ObjectId
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from docx import Document

import main


def sample_pdf(text="Education\nB.Tech Computer Science\n\nSkills\nPython; FastAPI\n\nProjects\nStudy assistant"):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (" + text.replace("\n", ") Tj 0 -20 Td (").encode() + b") Tj ET")
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    page[NameObject("/Contents")] = content
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def sample_docx():
    document = Document()
    document.add_heading("Education", level=1)
    document.add_paragraph("B.Tech Computer Science")
    document.add_heading("Skills", level=1)
    document.add_paragraph("Python; FastAPI")
    document.add_heading("Projects", level=1)
    document.add_paragraph("Study assistant")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def profile_payload(email="resume@example.com"):
    return {
        "personal_information": {"full_name": "Resume Student", "preferred_name": "Student", "email": email, "phone_number": "", "date_of_birth": "", "gender": "", "current_city": "Bengaluru", "country": "India"},
        "education": {"highest_education_level": "Undergraduate", "degree_course": "B.Tech", "specialization_branch": "Computer Science", "college_university": "Example University", "graduation_year": 2026, "current_education_status": "Currently studying", "year_semester": "4", "subjects": ["Python"]},
        "career_information": {"current_status": "Student", "target_job_role": "Developer", "preferred_industry": "Technology", "preferred_work_location": "Remote", "experience_level": "Entry level", "current_company": "", "current_job_role": ""},
        "skills": {"programming_languages": ["Python"], "technical_skills": [], "tools_technologies": [], "soft_skills": []},
        "career_goals": {"short_term_career_goal": "Find an internship", "long_term_career_goal": "", "target_role": "Developer", "target_companies": "", "areas_to_improve": "", "preferred_learning_style": [], "interests": "AI", "preferred_difficulty": "Medium"},
    }


class FakeCollection:
    def __init__(self):
        self.profiles = []
        self.versions = []

    async def find_one(self, query, projection=None, sort=None):
        items = self.versions if "version" in query or ("user_id" in query and self.versions) else self.profiles
        def matches_query(item):
            for key, value in query.items():
                actual = item.get(key)
                if key == "personal_information.email":
                    actual = item.get("personal_information", {}).get("email")
                if isinstance(value, dict) and "$gt" in value:
                    if actual is None or actual <= value["$gt"]:
                        return False
                elif actual != value:
                    return False
            return True
        matches = [item for item in items if matches_query(item)]
        if sort:
            matches.sort(key=lambda item: item.get(sort[0][0], 0), reverse=sort[0][1] < 0)
        return matches[0] if matches else None

    async def insert_one(self, document):
        document["_id"] = ObjectId()
        target = self.versions if "version" in document else self.profiles
        target.append(dict(document))

    async def update_one(self, query, update):
        for item in self.profiles:
            if all(item.get(key) == value for key, value in query.items()):
                for key, value in update.get("$set", {}).items():
                    if "." in key:
                        root, child, grandchild = key.split(".")
                        item.setdefault(root, {}).setdefault(child, {})[grandchild] = value
                    else:
                        item[key] = value
                return SimpleNamespace(modified_count=1, matched_count=1)
        return SimpleNamespace(modified_count=0, matched_count=0)


class FakeMongo:
    def __init__(self, collection):
        self.collection = collection

    def __getitem__(self, _database):
        return {"resume_versions": self.collection}


class ResumeTests(unittest.TestCase):
    def test_pdf_and_docx_extraction(self):
        pdf_text = main.extract_resume_text(".pdf", sample_pdf())
        docx_text = main.extract_resume_text(".docx", sample_docx())
        self.assertIn("Education", pdf_text)
        self.assertIn("Skills", docx_text)
        self.assertEqual(main.parse_resume_text(docx_text)["projects"], ["Study assistant"])
        inline_skills = main.parse_resume_text("SKILLS: Python, FastAPI, SQL\nProjects:\nStudy assistant")
        self.assertEqual(inline_skills["skills"], ["Python", "FastAPI", "SQL"])
        self.assertEqual(inline_skills["projects"], ["Study assistant"])
        with self.assertRaises(main.HTTPException):
            main.extract_resume_text(".pdf", sample_pdf(""))

    def test_invalid_and_empty_files_are_rejected(self):
        with self.assertRaises(main.HTTPException) as invalid:
            main.validate_resume_content("resume.pdf", b"not a pdf")
        self.assertEqual(invalid.exception.status_code, 422)
        with self.assertRaises(main.HTTPException) as empty:
            main.validate_resume_content("resume.docx", b"")
        self.assertEqual(empty.exception.status_code, 422)

    def test_persistence_replacement_and_version_history(self):
        collection = FakeCollection()
        client = TestClient(main.app)
        docx_content = sample_docx()
        with tempfile.TemporaryDirectory() as directory, patch.object(main, "profiles_collection", return_value=collection), patch.object(main, "mongo_client", FakeMongo(collection)), patch.object(main, "mongo_status", "connected"), patch.object(main, "UPLOAD_DIR", Path(directory)):
            form_data = {"profile_data": __import__("json").dumps(profile_payload()), "password": "SecureExamplePass123"}
            first = client.post("/api/users/profile", data=form_data, files={"resume": ("first.pdf", sample_pdf(), "application/pdf")})
            second = client.post("/api/users/profile", data=form_data, files={"resume": ("second.docx", docx_content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
            unchanged = client.post("/api/users/profile", data=form_data, files={"resume": ("renamed.docx", docx_content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        self.assertEqual(first.status_code, 201)
        self.assertEqual(second.status_code, 201)
        self.assertEqual(unchanged.status_code, 201)
        self.assertNotIn("session_token", first.json())
        self.assertIn("httponly", first.headers["set-cookie"].lower())
        self.assertNotEqual(collection.profiles[0]["password_hash"], "SecureExamplePass123")
        self.assertTrue(main.verify_password("SecureExamplePass123", collection.profiles[0]["password_hash"]))
        self.assertEqual(len(collection.versions), 2)
        self.assertEqual(collection.versions[0]["version"], 1)
        self.assertEqual(collection.versions[1]["version"], 2)
        self.assertEqual(collection.profiles[0]["resume_information"]["version"], 2)
        self.assertIn("Python", collection.profiles[0]["resume_information"]["extracted_information"]["skills"])
        self.assertIn("FastAPI", collection.profiles[0]["resume_information"]["extracted_information"]["skills"])
        self.assertEqual(collection.profiles[0]["resume_information"]["extracted_information"]["projects"], ["Study assistant"])
        self.assertIn("Python", collection.versions[0]["extracted_information"]["skills"])
        self.assertNotEqual(collection.versions[0]["stored_filename"], collection.versions[1]["stored_filename"])

    def test_existing_email_cannot_be_taken_over_without_authenticated_session(self):
        collection = FakeCollection()
        existing = profile_payload()
        collection.profiles.append({
            "user_id": "usr_existing",
            **existing,
            "password_hash": main.hash_password("ExistingSecurePassword123"),
        })
        with patch.object(main, "profiles_collection", return_value=collection), \
             patch.object(main, "mongo_client", FakeMongo(collection)), \
             patch.object(main, "mongo_status", "connected"), \
             patch.object(main, "profile_for_request", new=AsyncMock(return_value=None)):
            response = TestClient(main.app).post(
                "/api/users/profile",
                data={"profile_data": __import__("json").dumps(profile_payload()), "password": "AttackerPassword123"},
            )
        self.assertEqual(response.status_code, 409)
        self.assertEqual(collection.profiles[0]["user_id"], "usr_existing")

    def test_manual_resume_review_preserves_previous_version(self):
        user_id = "usr_review_owner"
        current_resume = {
            "version": 1,
            "uploaded": True,
            "stored_filename": "private-original.pdf",
            "extracted_information": {
                "education": ["B.Tech"],
                "skills": ["Python"],
                "projects": ["Study assistant"],
                "certifications": [],
                "internships": [],
                "work_experience": ["Developer"],
                "achievements": [],
            },
        }
        collection = FakeCollection()
        collection.profiles.append({"user_id": user_id, "resume_information": current_resume})
        collection.versions.append({"user_id": user_id, "version": 1, **current_resume})
        with patch.object(main, "mongo_client", FakeMongo(collection)), \
             patch.object(main, "mongo_status", "connected"), \
             patch.object(main, "profiles_collection", return_value=collection), \
             patch.object(main, "profile_for_request", new=AsyncMock(return_value={"user_id": user_id, "resume_information": current_resume})):
            response = TestClient(main.app).put(
                f"/api/users/{user_id}/resume",
                headers={"Authorization": "owner-token"},
                json={"extracted_information": {"skills": ["Python", "FastAPI"]}},
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["resume"]["version"], 2)
        self.assertEqual(len(collection.versions), 2)
        self.assertEqual(collection.versions[0]["extracted_information"]["skills"], ["Python"])
        self.assertEqual(collection.versions[1]["extracted_information"]["skills"], ["Python", "FastAPI"])
        self.assertEqual(collection.versions[1]["stored_filename"], "private-original.pdf")
        self.assertEqual(collection.versions[1]["extracted_information"]["education"], [])

    def test_resume_skill_update_requires_confirmation_and_creates_new_version(self):
        user_id = "usr_resume_owner"
        existing_resume = {
            "version": 1,
            "uploaded": True,
            "original_filename": "original.pdf",
            "stored_filename": "private-original.pdf",
            "file_sha256": "original-sha",
            "extracted_information": {
                "education": ["B.Tech"],
                "skills": ["FastAPI"],
                "projects": ["Study assistant"],
                "certifications": ["Existing certification"],
                "internships": [],
                "work_experience": ["Existing experience"],
                "achievements": [],
            },
        }
        collection = FakeCollection()
        collection.profiles.append({"user_id": user_id, "resume_information": existing_resume})
        original_version = {"user_id": user_id, "version": 1, **existing_resume}
        collection.versions.append(dict(original_version))
        qualification = {
            "qualified_skills": [{
                "skill": "Python",
                "assessment_id": "test-python",
                "score": 85,
                "completed_at": "2025-01-01T00:00:00+00:00",
            }],
        }
        with patch.object(main, "mongo_client", FakeMongo(collection)), \
             patch.object(main, "mongo_status", "connected"), \
             patch.object(main, "profiles_collection", return_value=collection), \
             patch.object(main, "profile_for_request", new=AsyncMock(side_effect=lambda request: {
                 "user_id": user_id,
                 "resume_information": collection.profiles[0]["resume_information"],
                 "career_information": {"target_job_role": "Python Developer"},
                 "career_goals": {},
                 "skills": {},
             })), \
             patch.object(main, "skill_path_data", new=AsyncMock(return_value=qualification)):
            client = TestClient(main.app)
            refused = client.post(
                f"/api/users/{user_id}/resume/qualified-skill?skill=Python",
                headers={"Authorization": "owner-token"},
                json={"confirmed": False},
            )
            self.assertEqual(refused.status_code, 400)
            self.assertEqual(len(collection.versions), 1)

            confirmed = client.post(
                f"/api/users/{user_id}/resume/qualified-skill?skill=Python",
                headers={"Authorization": "owner-token"},
                json={"confirmed": True},
            )
            self.assertEqual(confirmed.status_code, 200, confirmed.text)
            body = confirmed.json()
            self.assertTrue(body["added"])
            self.assertFalse(body["download_supported"])
            self.assertEqual(len(collection.versions), 2)
            self.assertEqual(collection.versions[0]["extracted_information"]["skills"], ["FastAPI"])
            self.assertEqual(collection.versions[1]["version"], 2)
            self.assertEqual(collection.versions[1]["prior_version"], 1)
            self.assertEqual(collection.versions[1]["stored_filename"], "private-original.pdf")
            self.assertEqual(collection.versions[1]["extracted_information"]["skills"], ["FastAPI", "Python"])
            self.assertEqual(collection.versions[1]["extracted_information"]["education"], ["B.Tech"])
            self.assertEqual(collection.versions[1]["extracted_information"]["projects"], ["Study assistant"])
            self.assertEqual(collection.versions[1]["extracted_information"]["work_experience"], ["Existing experience"])
            self.assertEqual(collection.versions[1]["skill_evidence"][0]["assessment_id"], "test-python")
            self.assertEqual(collection.profiles[0]["resume_information"]["version"], 2)
            self.assertEqual(collection.profiles[0]["resume_information"]["skill_evidence"][0]["assessment_score"], 85)

            duplicate = client.post(
                f"/api/users/{user_id}/resume/qualified-skill?skill=Python",
                headers={"Authorization": "owner-token"},
                json={"confirmed": True},
            )
            self.assertEqual(duplicate.status_code, 200)
            self.assertFalse(duplicate.json()["added"])
            self.assertEqual(len(collection.versions), 2)

    def test_resume_skill_update_rejects_ineligible_and_other_students(self):
        user_id = "usr_resume_owner"
        collection = FakeCollection()
        collection.profiles.append({"user_id": user_id, "resume_information": {"version": 1}})
        profile = {"user_id": user_id, "resume_information": {"version": 1}, "skills": {}}
        with patch.object(main, "mongo_client", FakeMongo(collection)), \
             patch.object(main, "mongo_status", "connected"), \
             patch.object(main, "profiles_collection", return_value=collection), \
             patch.object(main, "profile_for_request", new=AsyncMock(return_value=profile)), \
             patch.object(main, "skill_path_data", new=AsyncMock(return_value={"qualified_skills": []})):
            client = TestClient(main.app)
            ineligible = client.post(
                f"/api/users/{user_id}/resume/qualified-skill?skill=Python",
                headers={"Authorization": "owner-token"},
                json={"confirmed": True},
            )
            unauthorized = client.post(
                "/api/users/usr_other/resume/qualified-skill?skill=Python",
                headers={"Authorization": "owner-token"},
                json={"confirmed": True},
            )
        self.assertEqual(ineligible.status_code, 409)
        self.assertEqual(unauthorized.status_code, 403)
        self.assertEqual(collection.versions, [])

    def test_resume_endpoint_rejects_another_student(self):
        with patch.object(main, "profile_for_request", new=AsyncMock(return_value={"user_id": "usr_owner"})):
            response = TestClient(main.app).get("/api/users/usr_other/resume", headers={"Authorization": "Bearer token"})
        self.assertEqual(response.status_code, 403)


if __name__ == "__main__":
    unittest.main()
