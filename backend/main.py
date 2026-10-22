import json
import asyncio
import math
from io import BytesIO
import hashlib
import logging
import os
import re
import secrets
import time
import uuid
import zipfile
import xml.etree.ElementTree as ET
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode, urlparse
from urllib.request import Request as UrlRequest, urlopen
from email.utils import parsedate_to_datetime
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Literal, Optional

from fastapi import FastAPI, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from motor.motor_asyncio import AsyncIOMotorClient
from pymongo import ReturnDocument
from pymongo.errors import DuplicateKeyError
from pydantic import BaseModel, ConfigDict, EmailStr, Field, ValidationError, field_validator
from dotenv import load_dotenv
from google import genai
from google.genai import errors as genai_errors
from google.genai import types

BASE_DIR = Path(__file__).resolve().parent
FRONTEND_DIR = BASE_DIR.parent / "frontend"
ENV_FILE = BASE_DIR / ".env"
load_dotenv(dotenv_path=ENV_FILE, override=True)
MONGODB_URI = os.getenv("MONGODB_URI") or os.getenv("MONGODB_URL", "mongodb://127.0.0.1:27017/")
MONGODB_DATABASE = os.getenv("MONGODB_DATABASE", "krishna_ai_assistant")
MONGODB_COLLECTION = os.getenv("MONGODB_COLLECTION", "user_profiles")
UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", str(BASE_DIR / "uploads")))
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash").strip() or "gemini-2.5-flash"

def env_int(name: str, default: int, minimum: int, maximum: int) -> int:
    raw_value = os.getenv(name)
    if raw_value is None or not raw_value.strip():
        return default
    try:
        value = int(raw_value)
    except ValueError as exc:
        raise RuntimeError(f"{name} must be an integer.") from exc
    if not minimum <= value <= maximum:
        raise RuntimeError(f"{name} must be between {minimum} and {maximum}.")
    return value


def env_bool(name: str, default: bool = False) -> bool:
    raw_value = os.getenv(name)
    if raw_value is None or not raw_value.strip():
        return default
    normalized = raw_value.strip().lower()
    if normalized not in {"true", "false", "1", "0", "yes", "no"}:
        raise RuntimeError(f"{name} must be true or false.")
    return normalized in {"true", "1", "yes"}

GEMINI_MAX_OUTPUT_TOKENS = env_int("GEMINI_MAX_OUTPUT_TOKENS", 800, 1, 8192)
CHAT_MAX_MESSAGE_CHARS = env_int("CHAT_MAX_MESSAGE_CHARS", 4000, 1, 20000)
CHAT_MAX_HISTORY_MESSAGES = env_int("CHAT_MAX_HISTORY_MESSAGES", 20, 0, 50)
CHAT_MAX_CONTEXT_TOKENS = env_int("CHAT_MAX_CONTEXT_TOKENS", 16000, 1000, 1000000)
CHAT_MAX_REQUESTS = env_int("CHAT_MAX_REQUESTS", 20, 1, 10000)
CHAT_RATE_LIMIT_WINDOW = env_int("CHAT_RATE_LIMIT_WINDOW", 3600, 1, 86400)
CHAT_TIMEOUT_SECONDS = env_int("CHAT_TIMEOUT_SECONDS", 45, 5, 300)
AUTH_SESSION_TTL_SECONDS = env_int("AUTH_SESSION_TTL_SECONDS", 604800, 3600, 2592000)
AUTH_COOKIE_SECURE = env_bool("AUTH_COOKIE_SECURE", False)
GEMINI_MAX_RETRIES = env_int("GEMINI_MAX_RETRIES", 2, 0, 3)
GEMINI_RETRY_BASE_SECONDS = env_int("GEMINI_RETRY_BASE_SECONDS", 1, 1, 10)
TEST_MAX_OUTPUT_TOKENS = env_int("TEST_MAX_OUTPUT_TOKENS", 5000, 1000, 8192)
TEST_MAX_TOPIC_CHARS = env_int("TEST_MAX_TOPIC_CHARS", 300, 1, 2000)
TEST_MAX_LESSON_CHARS = env_int("TEST_MAX_LESSON_CHARS", 12000, 1000, 50000)
TEST_ACTIVE_TTL_SECONDS = env_int("TEST_ACTIVE_TTL_SECONDS", 86400, 300, 604800)
PASSING_SCORE_PERCENTAGE = env_int("PASSING_SCORE_PERCENTAGE", 70, 1, 100)
MIN_TOPIC_QUESTIONS = env_int("MIN_TOPIC_QUESTIONS", 2, 1, 50)
SKILL_LEVEL_DEVELOPING_MIN = env_int("SKILL_LEVEL_DEVELOPING_MIN", 50, 1, 99)
SKILL_LEVEL_PROFICIENT_MIN = env_int("SKILL_LEVEL_PROFICIENT_MIN", 70, 1, 100)
SKILL_LEVEL_STRONG_MIN = env_int("SKILL_LEVEL_STRONG_MIN", 85, 1, 100)
if not SKILL_LEVEL_DEVELOPING_MIN < SKILL_LEVEL_PROFICIENT_MIN < SKILL_LEVEL_STRONG_MIN:
    raise RuntimeError("Skill-level percentage thresholds must be strictly increasing.")
ENABLE_MOCK_QUIZ = env_bool("ENABLE_MOCK_QUIZ", False)
ADZUNA_APP_ID = os.getenv("ADZUNA_APP_ID", "").strip()
ADZUNA_APP_KEY = os.getenv("ADZUNA_APP_KEY", "").strip()
ADZUNA_COUNTRY = os.getenv("ADZUNA_COUNTRY", "in").strip().lower()
ADZUNA_ENABLED = env_bool("ADZUNA_ENABLED", False)
RESUME_MAX_BYTES = env_int("RESUME_MAX_BYTES", 10 * 1024 * 1024, 1, 50 * 1024 * 1024)
CHAT_INPUT_TOKEN_BUDGET = CHAT_MAX_CONTEXT_TOKENS - GEMINI_MAX_OUTPUT_TOKENS
if CHAT_INPUT_TOKEN_BUDGET < 1:
    raise RuntimeError("CHAT_MAX_CONTEXT_TOKENS must be greater than GEMINI_MAX_OUTPUT_TOKENS.")
SYSTEM_INSTRUCTION = (
    "You are Vaishu's AI Assistant, a warm and practical career companion. "
    "Give clear, actionable answers and ask a short follow-up question when useful."
)
gemini_client = genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None

logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"), format="%(asctime)s %(levelname)s %(name)s: %(message)s")
logger = logging.getLogger("vaishu-ai-assistant")


def redact_mongodb_secrets(message: str) -> str:
    """Remove credentials from driver errors before they reach logs or API responses."""
    return re.sub(r"(mongodb(?:\+srv)?://[^:]+:)[^@]+@", r"\1***@", message)


def explain_mongodb_error(exc: Exception) -> str:
    """Return a useful, password-safe diagnosis for health checks and logs."""
    message = redact_mongodb_secrets(f"{type(exc).__name__}: {exc}")
    lowered = message.lower()
    if "authentication failed" in lowered or "bad auth" in lowered:
        return "MongoDB authentication failed. Check the Atlas database username/password and URL-encode special password characters."
    if "serverselectiontimeout" in lowered or "timed out" in lowered:
        return "MongoDB connection timed out. Check the Atlas IP access list and network connectivity."
    if "name or service not known" in lowered or "nodename nor servname" in lowered or "getaddrinfo" in lowered:
        return "MongoDB hostname could not be resolved. Check MONGODB_URI."
    return message

app = FastAPI(title="Vaishu's AI Assistant API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip().rstrip("/") for origin in os.getenv(
        "FRONTEND_ORIGINS", "http://localhost:5500,http://127.0.0.1:5500"
    ).split(",") if origin.strip()],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
mongo_client: Optional[AsyncIOMotorClient] = None
mongo_status = "not_initialized"
mongo_error: Optional[str] = None
fallback_rate_limit: dict[str, list[float]] = {}
fallback_rate_limit_lock = asyncio.Lock()
fallback_active_tests: dict[str, dict] = {}
fallback_skill_verifications: dict[str, list[dict]] = {}
fallback_learning_plans: dict[str, list[dict]] = {}
fallback_topic_learning_progress: dict[str, dict[tuple[str, str], dict]] = {}

class PersonalInformation(BaseModel):
    full_name: str = Field(min_length=2)
    preferred_name: Optional[str] = None
    email: EmailStr
    phone_number: Optional[str] = None
    date_of_birth: Optional[str] = None
    gender: Optional[str] = None
    current_city: str = Field(min_length=2)
    country: str = Field(min_length=2)

class Education(BaseModel):
    highest_education_level: str
    degree_course: str
    specialization_branch: Optional[str] = None
    college_university: Optional[str] = None
    graduation_year: Optional[int] = Field(default=None, ge=1950, le=2100)
    current_education_status: str
    year_semester: Optional[str] = None
    subjects: list[str] = []

class CareerInformation(BaseModel):
    current_status: str
    target_job_role: str
    preferred_industry: Optional[str] = None
    preferred_work_location: Optional[str] = None
    experience_level: str
    current_company: Optional[str] = None
    current_job_role: Optional[str] = None

class Skills(BaseModel):
    programming_languages: list[str] = []
    technical_skills: list[str] = []
    tools_technologies: list[str] = []
    soft_skills: list[str] = []

class CareerGoals(BaseModel):
    short_term_career_goal: str
    long_term_career_goal: Optional[str] = None
    target_role: Optional[str] = None
    target_companies: Optional[str] = None
    areas_to_improve: Optional[str] = None
    preferred_learning_style: list[str] = []
    interests: Optional[str] = None
    preferred_difficulty: str = "Medium"

class Profile(BaseModel):
    model_config = ConfigDict(extra="forbid")
    personal_information: PersonalInformation
    education: Education
    career_information: CareerInformation
    skills: Skills
    career_goals: CareerGoals


class LoginRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    email: EmailStr
    password: str = Field(min_length=1, max_length=128)


class JobPreferencesRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    target_role: Optional[str] = Field(default=None, max_length=120)
    location: Optional[str] = Field(default=None, max_length=120)
    remote: Optional[bool] = None


class LearningCalendarRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    weekly_minutes: int = Field(ge=30, le=3000)
    start_date: str = Field(pattern=r"^\d{4}-\d{2}-\d{2}$")


class LearningCalendarRescheduleRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    task_id: str = Field(min_length=1, max_length=128)
    scheduled_date: str = Field(pattern=r"^\d{4}-\d{2}-\d{2}$")


class ResumeInformation(BaseModel):
    education: list[str] = Field(default_factory=list, max_length=30)
    skills: list[str] = Field(default_factory=list, max_length=50)
    projects: list[str] = Field(default_factory=list, max_length=30)
    certifications: list[str] = Field(default_factory=list, max_length=30)
    internships: list[str] = Field(default_factory=list, max_length=30)
    work_experience: list[str] = Field(default_factory=list, max_length=30)
    achievements: list[str] = Field(default_factory=list, max_length=30)


class ResumeReviewRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    extracted_information: ResumeInformation


class ResumeSkillUpdateRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    confirmed: bool


class LearningTopicTaskRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    skill: str = Field(min_length=1, max_length=120)
    topic: str = Field(min_length=1, max_length=200)
    task_id: str = Field(min_length=1, max_length=60)
    completed: bool


class SkillMapSelectionRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    target_role: str = Field(min_length=1, max_length=120)


class ChatMessage(BaseModel):
    model_config = ConfigDict(extra="forbid")
    role: Literal["user", "assistant"]
    content: str = Field(min_length=1, max_length=CHAT_MAX_MESSAGE_CHARS)

    @field_validator("content")
    @classmethod
    def non_blank_content(cls, value: str) -> str:
        value = value.strip()
        if not value:
            raise ValueError("Message content cannot be blank.")
        return value


class ChatRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    message: str = Field(min_length=1, max_length=CHAT_MAX_MESSAGE_CHARS)
    history: list[ChatMessage] = Field(default_factory=list, max_length=CHAT_MAX_HISTORY_MESSAGES)

    @field_validator("message")
    @classmethod
    def non_blank_message(cls, value: str) -> str:
        value = value.strip()
        if not value:
            raise ValueError("Message cannot be blank.")
        return value


OptionKey = Literal["A", "B", "C", "D"]
Difficulty = Literal["Easy", "Medium", "Hard", "Beginner", "Intermediate", "Advanced"]
QuestionType = Literal["mcq"]


class TestGenerateRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    topic: Optional[str] = Field(default=None, max_length=TEST_MAX_TOPIC_CHARS)
    subject: Optional[str] = Field(default=None, max_length=100)
    difficulty: Difficulty = "Medium"
    question_count: Literal[5, 10, 15] = 5
    question_type: QuestionType = "mcq"
    history: list[ChatMessage] = Field(default_factory=list, max_length=CHAT_MAX_HISTORY_MESSAGES)
    use_mock: bool = False

    @field_validator("topic")
    @classmethod
    def normalize_topic(cls, value: Optional[str]) -> Optional[str]:
        if value is None:
            return None
        value = value.strip()
        return value or None

    @field_validator("subject")
    @classmethod
    def normalize_subject(cls, value: Optional[str]) -> Optional[str]:
        if value is None:
            return None
        value = value.strip()
        return value or None


class TestSubmitRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    test_id: str = Field(min_length=8, max_length=80)
    answers: dict[str, OptionKey] = Field(default_factory=dict, max_length=15)


class LearningPlanProgressRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    task_index: int = Field(ge=0, le=30)
    completed: bool


class GeneratedQuestion(BaseModel):
    question_id: str = Field(min_length=1, max_length=40)
    prompt: str = Field(min_length=10, max_length=500)
    option_a: str = Field(min_length=1, max_length=250)
    option_b: str = Field(min_length=1, max_length=250)
    option_c: str = Field(min_length=1, max_length=250)
    option_d: str = Field(min_length=1, max_length=250)
    correct_answer: OptionKey
    explanation: str = Field(min_length=1, max_length=600)
    topic: str = Field(min_length=1, max_length=200)
    question_type: QuestionType = "mcq"
    difficulty: Difficulty = "Medium"
    concept: Optional[str] = Field(default=None, max_length=200)


class GeneratedTest(BaseModel):
    questions: list[GeneratedQuestion]

@app.on_event("startup")
async def connect_to_mongodb():
    global mongo_client, mongo_status, mongo_error
    if not MONGODB_URI:
        mongo_status = "configuration_error"
        mongo_error = "MONGODB_URI is missing from the environment."
        logger.error(mongo_error)
        return
    logger.info("Connecting to MongoDB database=%s", MONGODB_DATABASE)
    mongo_client = AsyncIOMotorClient(
        MONGODB_URI,
        serverSelectionTimeoutMS=10000,
        connectTimeoutMS=10000,
        socketTimeoutMS=20000,
        appname="VaishuAI",
    )
    try:
        await mongo_client.admin.command("ping")
        await mongo_client[MONGODB_DATABASE][MONGODB_COLLECTION].create_index(
            "personal_information.email", unique=True, name="unique_profile_email"
        )
        await mongo_client[MONGODB_DATABASE]["chat_rate_limits"].create_index(
            "created_at", expireAfterSeconds=max(CHAT_RATE_LIMIT_WINDOW * 2, 3600), name="chat_rate_limit_ttl"
        )
        await mongo_client[MONGODB_DATABASE]["learning_tests"].create_index(
            "expires_at", expireAfterSeconds=0, name="learning_test_expiry"
        )
        await mongo_client[MONGODB_DATABASE]["resume_versions"].create_index(
            [("user_id", 1), ("version", -1)], name="resume_user_version"
        )
        await mongo_client[MONGODB_DATABASE]["assessment_attempts"].create_index(
            [("user_id", 1), ("submitted_at", -1)], name="assessment_user_submitted"
        )
        await mongo_client[MONGODB_DATABASE]["skill_verifications"].create_index(
            [("user_id", 1), ("skill_key", 1)], unique=True, name="verified_user_skill_key"
        )
        await mongo_client[MONGODB_DATABASE]["learning_plans"].create_index(
            [("user_id", 1), ("created_at", -1)], name="learning_plan_user_created"
        )
        await mongo_client[MONGODB_DATABASE]["topic_learning_progress"].create_index(
            [("user_id", 1), ("skill_key", 1), ("topic_key", 1)],
            unique=True,
            name="topic_progress_user_skill_topic",
        )
        await mongo_client[MONGODB_DATABASE]["job_recommendations"].create_index(
            [("user_id", 1), ("day_key", 1)], unique=True, name="job_user_day"
        )
        await mongo_client[MONGODB_DATABASE]["job_preferences"].create_index(
            "user_id", unique=True, name="job_preferences_user"
        )
        await mongo_client[MONGODB_DATABASE]["job_provider_cache"].create_index(
            "cache_key", unique=True, name="job_provider_query_cache"
        )
        await mongo_client[MONGODB_DATABASE]["learning_calendars"].create_index(
            "user_id", unique=True, name="learning_calendar_user"
        )
        mongo_status = "connected"
        mongo_error = None
        logger.info("MongoDB connected: database=%s collection=%s", MONGODB_DATABASE, MONGODB_COLLECTION)
    except Exception as exc:
        mongo_status = "unavailable"
        mongo_error = explain_mongodb_error(exc)
        logger.error("MongoDB connection failed: %s", mongo_error)

@app.on_event("shutdown")
async def close_mongodb():
    if mongo_client:
        mongo_client.close()


def profiles_collection():
    if mongo_client is None:
        raise HTTPException(status_code=503, detail="Database connection is not ready.")
    return mongo_client[MONGODB_DATABASE][MONGODB_COLLECTION]


def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    digest = hashlib.scrypt(password.encode("utf-8"), salt=salt, n=16384, r=8, p=1, dklen=64)
    return f"scrypt$16384$8$1${salt.hex()}${digest.hex()}"


def verify_password(password: str, encoded: Optional[str]) -> bool:
    if not encoded:
        return False
    try:
        algorithm, n, r, p, salt_hex, digest_hex = encoded.split("$")
        if algorithm != "scrypt":
            return False
        digest = hashlib.scrypt(
            password.encode("utf-8"),
            salt=bytes.fromhex(salt_hex),
            n=int(n),
            r=int(r),
            p=int(p),
            dklen=len(bytes.fromhex(digest_hex)),
        )
        return secrets.compare_digest(digest.hex(), digest_hex)
    except (ValueError, TypeError, MemoryError):
        return False


def set_session_cookie(response: Response, token: str) -> None:
    response.set_cookie(
        "vaishu_session",
        token,
        max_age=AUTH_SESSION_TTL_SECONDS,
        httponly=True,
        secure=AUTH_COOKIE_SECURE,
        samesite="lax",
        path="/",
    )


def clear_session_cookie(response: Response) -> None:
    response.delete_cookie("vaishu_session", path="/", httponly=True, secure=AUTH_COOKIE_SECURE, samesite="lax")


async def profile_for_request(request: Request) -> Optional[dict]:
    authorization = request.headers.get("authorization", "")
    tokens = []
    if authorization.lower().startswith("bearer "):
        tokens.append(authorization[7:].strip())
    cookie_token = request.cookies.get("vaishu_session", "")
    if cookie_token and cookie_token not in tokens:
        tokens.append(cookie_token)
    if not tokens:
        return None
    if mongo_status != "connected":
        raise HTTPException(status_code=503, detail="Student session verification is unavailable because the profile database is disconnected.")
    collection = profiles_collection()
    try:
        for token in tokens:
            token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
            profile = await collection.find_one(
                {"auth_token_hash": token_hash, "auth_token_expires_at": {"$gt": datetime.now(timezone.utc)}},
                {"_id": 0},
            )
            if profile:
                return profile
        return None
    except Exception as exc:
        logger.error("Unable to verify a student session")
        raise HTTPException(status_code=503, detail="Student session verification is temporarily unavailable.") from exc


async def require_profile(request: Request) -> dict:
    profile = await profile_for_request(request)
    if profile is None:
        raise HTTPException(status_code=401, detail="Your student session is missing or expired. Please return to onboarding and submit your profile again.")
    return profile


def profile_context(profile: Optional[dict]) -> str:
    if not profile:
        return ""
    personal = profile.get("personal_information", {})
    education = profile.get("education", {})
    skills = profile.get("skills", {})
    goals = profile.get("career_goals", {})
    resume = profile.get("resume_information", {}).get("extracted_information", {})
    name = personal.get("preferred_name") or personal.get("full_name") or "the student"
    subjects = ", ".join(education.get("subjects") or [])
    skills_text = ", ".join(
        item
        for key in ("programming_languages", "technical_skills", "tools_technologies", "soft_skills")
        for item in (skills.get(key) or [])
    )
    details = [
        f"Student name: {name}",
        f"Education: {education.get('highest_education_level', '')} {education.get('degree_course', '')}".strip(),
        f"Branch: {education.get('specialization_branch', '')}" if education.get("specialization_branch") else "",
        f"Year/semester: {education.get('year_semester', '')}" if education.get("year_semester") else "",
        f"Subjects: {subjects}" if subjects else "",
        f"Skills: {skills_text}" if skills_text else "",
        f"Interests: {goals.get('interests', '')}" if goals.get("interests") else "",
        f"Learning goals: {goals.get('short_term_career_goal', '')}" if goals.get("short_term_career_goal") else "",
        f"Preferred difficulty: {goals.get('preferred_difficulty', 'Medium')}",
        f"Resume education: {', '.join(resume.get('education') or [])}" if resume.get("education") else "",
        f"Resume skills: {', '.join(resume.get('skills') or [])}" if resume.get("skills") else "",
        f"Resume projects: {', '.join(resume.get('projects') or [])}" if resume.get("projects") else "",
        f"Resume experience: {', '.join(resume.get('work_experience') or [])}" if resume.get("work_experience") else "",
    ]
    return "Student profile context (use when relevant; do not mention these details unless helpful):\n" + "\n".join(item for item in details if item)


def serialize_document(value):
    """Convert MongoDB-specific values into JSON-safe values."""
    if isinstance(value, dict):
        return {key: serialize_document(item) for key, item in value.items() if key not in {"_id", "auth_token_hash", "auth_token_expires_at", "password_hash", "stored_filename"}}
    if isinstance(value, list):
        return [serialize_document(item) for item in value]
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value) if value.__class__.__name__ == "ObjectId" else value


RESUME_SECTIONS = {
    "education": ("education", "academic background", "academic qualifications"),
    "skills": ("skills", "technical skills", "core skills", "technologies"),
    "projects": ("projects", "personal projects", "academic projects"),
    "certifications": ("certifications", "certificates"),
    "internships": ("internships", "internship experience"),
    "work_experience": ("experience", "work experience", "professional experience", "employment"),
    "achievements": ("achievements", "awards", "accomplishments"),
}


def clean_resume_line(line: str) -> str:
    return re.sub(r"^\s*(?:[-*•▪◦]|\d+[.)])\s*", "", line).strip()


def resume_section_name(line: str) -> Optional[str]:
    normalized = re.sub(r"[^a-z ]", "", line.casefold()).strip()
    if len(normalized) > 60:
        return None
    for section, aliases in RESUME_SECTIONS.items():
        if normalized in aliases:
            return section
    return None


def parse_resume_text(text: str) -> dict[str, list[str]]:
    """Extract explicitly labelled resume sections without inventing content."""
    extracted = {section: [] for section in RESUME_SECTIONS}
    current: Optional[str] = None
    for raw_line in text.splitlines():
        line = clean_resume_line(raw_line)
        if not line:
            continue
        section = None
        inline_content = ""
        for name, aliases in RESUME_SECTIONS.items():
            alias_pattern = "|".join(re.escape(alias) for alias in sorted(aliases, key=len, reverse=True))
            match = re.match(rf"^\s*(?:{alias_pattern})\s*(?::|-)?\s*(.*)$", line, flags=re.IGNORECASE)
            if match:
                section = name
                inline_content = match.group(1).strip()
                break
        if section:
            current = section
            line = inline_content
            if not line:
                continue
        if current and len(line) >= 2:
            delimiter = r"\s*[|;,]\s*" if current == "skills" else r"\s*[|;]\s*"
            values = [part.strip() for part in re.split(delimiter, line) if part.strip()]
            for value in values:
                if value not in extracted[current] and len(extracted[current]) < 30:
                    extracted[current].append(value[:500])
    return extracted


def validate_resume_content(filename: str, content: bytes) -> str:
    extension = Path(filename or "").suffix.lower()
    if extension not in {".pdf", ".docx"}:
        raise HTTPException(status_code=400, detail="Resume must be a PDF or DOCX file.")
    if not content:
        raise HTTPException(status_code=422, detail="The uploaded resume is empty.")
    if len(content) > RESUME_MAX_BYTES:
        raise HTTPException(status_code=413, detail=f"Resume must be smaller than {RESUME_MAX_BYTES // (1024 * 1024)} MB.")
    if extension == ".pdf" and not content.startswith(b"%PDF-"):
        raise HTTPException(status_code=422, detail="The uploaded file is not a valid PDF.")
    if extension == ".docx":
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("Missing Word document content")
                ET.fromstring(archive.read("word/document.xml"))
        except (ValueError, ET.ParseError, zipfile.BadZipFile) as exc:
            raise HTTPException(status_code=422, detail="The uploaded file is not a valid DOCX document.") from exc
    return extension


def extract_resume_text(extension: str, content: bytes) -> str:
    try:
        if extension == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            from docx import Document
            document = Document(BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    except Exception as exc:
        logger.warning("Resume extraction failed: type=%s error=%s", extension, type(exc).__name__)
        raise HTTPException(status_code=422, detail="The resume could not be read. Please upload a valid text-based PDF or DOCX file.") from exc
    text = re.sub(r"[ \t]+", " ", text).strip()
    if not text:
        raise HTTPException(status_code=422, detail="The resume contains no extractable text. Please upload a text-based PDF or DOCX file.")
    return text[:200000]


async def next_resume_version(collection, user_id: str) -> int:
    latest = await collection.find_one({"user_id": user_id}, {"version": 1}, sort=[("version", -1)])
    return int((latest or {}).get("version", 0)) + 1


def public_resume_information(info: dict) -> dict:
    return {key: value for key, value in info.items() if key not in {"_id", "stored_filename", "user_id"}}

@app.get("/health")
@app.get("/api/health")
async def health(response: Response):
    global mongo_status, mongo_error
    try:
        if mongo_client is None:
            raise RuntimeError("MongoDB client is not initialized")
        await mongo_client.admin.command("ping")
        mongo_status = "connected"
        mongo_error = None
    except Exception as exc:
        mongo_status = "unavailable"
        mongo_error = explain_mongodb_error(exc)
        logger.error("MongoDB health check failed: %s", mongo_error)
    if mongo_status != "connected":
        response.status_code = 503
        return {"status": "unavailable", "database": mongo_status}
    return {"status": "ok", "database": "connected", "service": "Vaishu's AI Assistant"}


def build_chat_contents(history: list[ChatMessage], message: str) -> list[types.Content]:
    contents = [
        types.Content(
            role="user" if item.role == "user" else "model",
            parts=[types.Part.from_text(text=item.content)],
        )
        for item in history
    ]
    contents.append(types.Content(role="user", parts=[types.Part.from_text(text=message)]))
    return contents


def estimated_content_tokens(contents: list[types.Content]) -> int:
    """Fallback estimate only; provider counts are preferred and returned when available."""
    characters = sum(len(part.text or "") for content in contents for part in (content.parts or []))
    return max(1, (characters + 3) // 4)


async def count_content_tokens(contents: list[types.Content]) -> tuple[Optional[int], bool]:
    if gemini_client is None:
        return None, True
    try:
        result = await asyncio.wait_for(
            asyncio.to_thread(
                gemini_client.models.count_tokens,
                model=GEMINI_MODEL,
                contents=[types.Content(role="user", parts=[types.Part.from_text(text=SYSTEM_INSTRUCTION)]), *contents],
            ),
            timeout=CHAT_TIMEOUT_SECONDS,
        )
        token_count = getattr(result, "total_tokens", None)
        if isinstance(token_count, int) and token_count > 0:
            return token_count, False
    except Exception:
        logger.warning("Gemini token count unavailable; using an internal estimate")
    return None, True


async def fit_history(history: list[ChatMessage], message: str) -> tuple[list[ChatMessage], list[types.Content], int, bool]:
    selected = list(history[-CHAT_MAX_HISTORY_MESSAGES:])
    while selected and selected[0].role == "assistant":
        selected.pop(0)
    contents = build_chat_contents(selected, message)
    provider_tokens, estimated = await count_content_tokens(contents)
    token_count = provider_tokens or estimated_content_tokens(contents)
    while selected and token_count > CHAT_INPUT_TOKEN_BUDGET:
        selected.pop(0)
        while selected and selected[0].role == "assistant":
            selected.pop(0)
        contents = build_chat_contents(selected, message)
        provider_tokens, estimated = await count_content_tokens(contents)
        token_count = provider_tokens or estimated_content_tokens(contents)
    if token_count > CHAT_INPUT_TOKEN_BUDGET:
        raise HTTPException(status_code=413, detail="That message and conversation are too large. Please start a shorter conversation.")
    return selected, contents, token_count, estimated


def session_key(request: Request) -> str:
    supplied = request.headers.get("authorization", "").strip()
    if supplied.lower().startswith("bearer "):
        identity = hashlib.sha256(supplied[7:].strip().encode("utf-8")).hexdigest()
    else:
        supplied = request.headers.get("x-chat-session", "").strip()
        identity = supplied[:128] if supplied else (request.client.host if request.client else "unknown")
    return hashlib.sha256(identity.encode("utf-8")).hexdigest()


async def check_rate_limit(request: Request) -> tuple[bool, int]:
    now = int(time.time())
    window_start = now // CHAT_RATE_LIMIT_WINDOW
    key = f"{session_key(request)}:{window_start}"
    document_id = hashlib.sha256(key.encode("utf-8")).hexdigest()
    if mongo_client is not None and mongo_status == "connected":
        try:
            document = await mongo_client[MONGODB_DATABASE]["chat_rate_limits"].find_one_and_update(
                {"_id": document_id},
                {"$inc": {"count": 1}, "$setOnInsert": {"created_at": datetime.now(timezone.utc), "window_start": window_start}},
                upsert=True,
                return_document=ReturnDocument.AFTER,
            )
            count = int(document.get("count", 0))
            return count <= CHAT_MAX_REQUESTS, max(0, CHAT_MAX_REQUESTS - count)
        except Exception:
            logger.warning("MongoDB rate limiter unavailable; using process-local fallback")
    identity = session_key(request)
    async with fallback_rate_limit_lock:
        cutoff = now - CHAT_RATE_LIMIT_WINDOW
        timestamps = [timestamp for timestamp in fallback_rate_limit.get(identity, []) if timestamp > cutoff]
        timestamps.append(float(now))
        fallback_rate_limit[identity] = timestamps
        return len(timestamps) <= CHAT_MAX_REQUESTS, max(0, CHAT_MAX_REQUESTS - len(timestamps))


def response_usage(response) -> Optional[dict[str, int]]:
    metadata = getattr(response, "usage_metadata", None)
    if metadata is None:
        return None
    fields = {
        "input_tokens": getattr(metadata, "prompt_token_count", None),
        "output_tokens": getattr(metadata, "candidates_token_count", None),
        "total_tokens": getattr(metadata, "total_token_count", None),
    }
    usage = {key: value for key, value in fields.items() if isinstance(value, int)}
    return usage or None


@app.post("/api/chat")
async def chat(request: ChatRequest, http_request: Request):
    if gemini_client is None:
        raise HTTPException(status_code=503, detail="Gemini is not configured. Set GEMINI_API_KEY in the backend environment.")
    profile = await require_profile(http_request)
    allowed, remaining = await check_rate_limit(http_request)
    if not allowed:
        raise HTTPException(status_code=429, detail="You have reached the chat request limit. Please try again later.", headers={"Retry-After": str(CHAT_RATE_LIMIT_WINDOW)})
    _, contents, input_tokens, estimated = await fit_history(request.history, request.message)
    student_context = profile_context(profile)
    system_instruction = "\n\n".join(part for part in (SYSTEM_INSTRUCTION, student_context) if part)
    try:
        response = await asyncio.wait_for(
            asyncio.to_thread(
                gemini_client.models.generate_content,
                model=GEMINI_MODEL,
                contents=contents,
                config=types.GenerateContentConfig(system_instruction=system_instruction, temperature=0.7, max_output_tokens=GEMINI_MAX_OUTPUT_TOKENS),
            ),
            timeout=CHAT_TIMEOUT_SECONDS,
        )
    except asyncio.TimeoutError as exc:
        logger.warning("Gemini request timed out after %ss", CHAT_TIMEOUT_SECONDS)
        raise HTTPException(status_code=504, detail="The assistant took too long to respond. Please try again.") from exc
    except Exception as exc:
        error_text = str(exc).lower()
        if "429" in error_text or "quota" in error_text or "resource exhausted" in error_text:
            raise HTTPException(status_code=429, detail="Gemini is temporarily rate-limited. Please try again later.") from exc
        logger.warning("Gemini request failed: %s", type(exc).__name__)
        raise HTTPException(status_code=502, detail="The assistant is temporarily unavailable. Please try again.") from exc
    answer = (response.text or "").strip()
    if not answer:
        raise HTTPException(status_code=502, detail="The assistant returned an empty response. Please try again.")
    return {"response": answer, "model": GEMINI_MODEL, "usage": response_usage(response), "preflight_input_tokens": input_tokens, "preflight_input_tokens_estimated": estimated, "rate_limit_remaining": remaining}


def test_question_for_storage(question: GeneratedQuestion) -> dict:
    return {
        "question_id": question.question_id,
        "prompt": question.prompt,
        "options": {"A": question.option_a, "B": question.option_b, "C": question.option_c, "D": question.option_d},
        "correct_answer": question.correct_answer,
        "explanation": question.explanation,
        "topic": question.topic,
        "question_type": question.question_type,
        "difficulty": question.difficulty,
        "concept": question.concept or question.topic,
    }


def public_test(test_document: dict) -> dict:
    return {
        "test_id": test_document["test_id"],
        "topic": test_document["topic"],
        "difficulty": test_document["difficulty"],
        "question_count": test_document["question_count"],
        "subject": test_document.get("subject"),
        "passing_percentage": PASSING_SCORE_PERCENTAGE,
        "source": test_document.get("source", "gemini"),
        "questions": [
            {"question_id": item["question_id"], "prompt": item["prompt"], "options": item["options"], "topic": item["topic"], "question_type": item.get("question_type", "mcq"), "difficulty": item.get("difficulty", test_document["difficulty"]), "concept": item.get("concept", item["topic"])}
            for item in test_document["questions"]
        ],
    }


def mock_test_payload(request: TestGenerateRequest) -> GeneratedTest:
    """Development-only sample data for testing the UI while Gemini is unavailable."""
    topic = request.topic or request.subject or "the selected lesson"
    templates = [
        (f"Which action best demonstrates understanding of {topic}?", "Apply the idea to a new example", "Memorize one definition", "Skip all practice", "Copy an answer", "A", "Applying the idea to a new example checks understanding."),
        (f"When studying {topic}, what is the most useful first step?", "Identify the core concept", "Ignore the requirements", "Choose an answer randomly", "Avoid examples", "A", "Identifying the core concept creates a foundation for practice."),
        (f"Which result would show progress with {topic}?", "Explaining the idea and solving a new problem", "Repeating a phrase without context", "Avoiding feedback", "Only reading the title", "A", "Transfer to a new problem is stronger evidence than recall alone."),
        (f"What is a good way to check an assumption about {topic}?", "Test it with a small example", "Treat it as always true", "Delete the example", "Avoid asking questions", "A", "A small example can reveal whether an assumption holds."),
        (f"Which study habit best supports long-term recall of {topic}?", "Spaced practice with retrieval", "One long session only", "Never reviewing mistakes", "Reading without attempting", "A", "Spaced retrieval strengthens recall and reveals gaps."),
        (f"If a solution involving {topic} fails, what should you do first?", "Inspect the result and isolate the cause", "Change everything at once", "Assume the tool is broken", "Submit without checking", "A", "Isolating the cause makes debugging focused and measurable."),
        (f"Which explanation of {topic} is strongest?", "One that connects the rule to an example", "One with no example", "One copied without understanding", "One that avoids the question", "A", "A connected example demonstrates both the rule and its use."),
        (f"How can you compare two approaches to {topic}?", "Use the same criteria and test case", "Change the criteria for each", "Compare only their names", "Avoid measuring results", "A", "Consistent criteria and test cases make comparisons meaningful."),
        (f"What should a learner do after finding a mistake in {topic}?", "Record the cause and retry the problem", "Hide the mistake", "Stop practicing", "Repeat the same step unchanged", "A", "Recording the cause turns a mistake into a learning opportunity."),
        (f"Which question encourages deeper thinking about {topic}?", "When would this approach not work?", "What is the first word?", "Can I skip the example?", "Is there only one phrase?", "A", "Considering limitations encourages transfer and deeper understanding."),
        (f"What makes practice for {topic} realistic?", "Using a problem similar to a real task", "Changing no inputs", "Avoiding constraints", "Only copying notes", "A", "Realistic practice connects learning to useful decisions."),
        (f"How should you respond to unclear feedback on {topic}?", "Ask for a specific example", "Ignore it", "Guess the issue", "Delete your work", "A", "A specific example makes feedback actionable."),
        (f"Which approach helps explain {topic} to a beginner?", "Use plain language and a small example", "Use unexplained jargon", "List unrelated facts", "Skip the main idea", "A", "Plain language and examples make a concept accessible."),
        (f"What is a useful next step after learning {topic}?", "Build a small project using it", "Never applying it", "Only changing fonts", "Avoiding questions", "A", "A small project provides practice and evidence of understanding."),
        (f"How can you tell whether you truly understand {topic}?", "Solve or explain it without copying", "Recognize the title only", "Read it once", "Avoid testing yourself", "A", "Independent explanation or application is a stronger check of understanding."),
    ]
    questions = []
    for index, (prompt, option_a, option_b, option_c, option_d, correct, explanation) in enumerate(templates[:request.question_count], start=1):
        questions.append(GeneratedQuestion(
            question_id=f"q{index}", prompt=prompt, option_a=option_a, option_b=option_b,
            option_c=option_c, option_d=option_d, correct_answer=correct,
            explanation=f"[Sample data] {explanation}", topic=topic[:200], question_type="mcq",
            difficulty=request.difficulty, concept=topic[:200],
        ))
    return GeneratedTest(questions=questions)


def test_prompt(request: TestGenerateRequest, lesson_context: str) -> str:
    return f"""Create a high-quality multiple-choice learning test.

Topic supplied by learner: {request.topic or '(infer from the lesson)'}
Subject: {request.subject or '(infer from topic and student profile)'}
Difficulty: {request.difficulty}
Number of questions: {request.question_count}
Question type: {request.question_type}

Lesson and conversation context:
{lesson_context}

Rules:
- Test understanding and application of this lesson, not trivia or unrelated general knowledge.
- Use only the lesson/context above. If a topic was supplied, use it to focus the test.
- Create exactly {request.question_count} distinct questions with exactly four plausible options.
- Make the options unambiguous and ensure exactly one correct answer per question.
- Use question_id values q1, q2, q3, etc. in order.
- Include a short explanation of why the correct answer is right, a specific topic label, a concept tag, question_type, and difficulty.
- Every question must be directly relevant to the requested subject/topic or the saved student profile skills.
- Return only the requested structured JSON object.
"""


def sanitize_gemini_error_message(exc: Exception) -> str:
    """Keep provider diagnostics useful while removing credentials and secrets."""
    message = str(getattr(exc, "message", None) or exc)
    message = re.sub(r"(?i)(api[_ -]?key|authorization|bearer)\s*[:=]\s*[^\s,;]+", r"\1=[REDACTED]", message)
    message = re.sub(r"AIza[0-9A-Za-z_-]{20,}", "[REDACTED]", message)
    return message[:500]


def classify_gemini_error(exc: Exception) -> tuple[str, Optional[int], Optional[str], str]:
    """Classify SDK/provider failures without relying only on string matching."""
    code = getattr(exc, "code", None)
    status = getattr(exc, "status", None)
    message = sanitize_gemini_error_message(exc)
    lowered = f"{status or ''} {message}".lower()
    if code in {401, 403} or any(marker in lowered for marker in ("invalid api key", "api key not valid", "unauthenticated", "permission denied")):
        category = "invalid_key"
    elif code == 429 or "429" in lowered:
        category = "quota_exhausted" if any(marker in lowered for marker in ("quota", "resource exhausted", "daily limit", "per minute limit")) else "transient_rate_limit"
    elif code in {500, 502, 503, 504} or isinstance(exc, genai_errors.ServerError):
        category = "temporary_server_error"
    else:
        category = "provider_error"
    return category, code, status, message


def gemini_retry_delay(exc: Exception, attempt: int) -> float:
    """Honor Retry-After/retryDelay, bounded by a small safety cap."""
    response = getattr(exc, "response", None)
    headers = getattr(response, "headers", {}) or {}
    retry_after = headers.get("retry-after") if hasattr(headers, "get") else None
    delay: Optional[float] = None
    if retry_after:
        try:
            delay = float(retry_after)
        except (TypeError, ValueError):
            try:
                retry_at = parsedate_to_datetime(retry_after)
                if retry_at.tzinfo is None:
                    retry_at = retry_at.replace(tzinfo=timezone.utc)
                delay = (retry_at - datetime.now(timezone.utc)).total_seconds()
            except (TypeError, ValueError, OverflowError):
                delay = None
    if delay is None:
        match = re.search(r"(?:retryDelay|retry_after|retry-after)[^0-9]*(\d+(?:\.\d+)?)\s*(h|m|s)?", str(getattr(exc, "details", "")), re.IGNORECASE)
        if match:
            units = {"h": 3600, "m": 60, "s": 1}
            delay = float(match.group(1)) * units.get((match.group(2) or "s").lower(), 1)
        else:
            delay = GEMINI_RETRY_BASE_SECONDS * (2 ** attempt)
    return max(0.1, min(delay, 10.0))


async def generate_test_payload(prompt: str, expected_count: int) -> GeneratedTest:
    if gemini_client is None:
        raise HTTPException(status_code=503, detail="Gemini is not configured. Set GEMINI_API_KEY in the backend environment.")
    config = types.GenerateContentConfig(
        response_mime_type="application/json",
        response_schema=GeneratedTest,
        temperature=0.4,
        max_output_tokens=TEST_MAX_OUTPUT_TOKENS,
    )
    for attempt in range(GEMINI_MAX_RETRIES + 1):
        try:
            response = await asyncio.wait_for(
                asyncio.to_thread(
                    gemini_client.models.generate_content,
                    model=GEMINI_MODEL,
                    contents=prompt,
                    config=config,
                ),
                timeout=CHAT_TIMEOUT_SECONDS,
            )
            break
        except asyncio.TimeoutError as exc:
            logger.warning("Gemini quiz generation timed out: model=%s attempt=%d/%d", GEMINI_MODEL, attempt + 1, GEMINI_MAX_RETRIES + 1)
            raise HTTPException(status_code=504, detail="Generating the test took too long. Please try again.") from exc
        except Exception as exc:
            category, code, status, provider_message = classify_gemini_error(exc)
            logger.warning(
                "Gemini quiz generation failed: category=%s code=%s status=%s attempt=%d/%d message=%s",
                category,
                code or "unknown",
                status or "unknown",
                attempt + 1,
                GEMINI_MAX_RETRIES + 1,
                provider_message,
            )
            if category == "transient_rate_limit" and attempt < GEMINI_MAX_RETRIES:
                delay = gemini_retry_delay(exc, attempt)
                logger.info("Retrying Gemini quiz generation in %.1f seconds", delay)
                await asyncio.sleep(delay)
                continue
            if category == "invalid_key":
                raise HTTPException(status_code=503, detail="Gemini authentication failed. Check GEMINI_API_KEY in the backend environment.") from exc
            if category == "quota_exhausted":
                raise HTTPException(status_code=429, detail="Gemini quota is exhausted. Please try again after the quota resets.") from exc
            if category == "transient_rate_limit":
                raise HTTPException(status_code=429, detail="Gemini is temporarily rate-limited. Please retry in a moment.") from exc
            if category == "temporary_server_error":
                raise HTTPException(status_code=503, detail="Gemini is temporarily unavailable. Please try again shortly.") from exc
            raise HTTPException(status_code=502, detail="The test could not be generated by Gemini. Please try again.") from exc
    try:
        parsed = getattr(response, "parsed", None)
        payload = parsed if isinstance(parsed, GeneratedTest) else GeneratedTest.model_validate_json(response.text or "")
    except (ValidationError, ValueError, json.JSONDecodeError) as exc:
        logger.warning("Gemini returned invalid test structure")
        raise HTTPException(status_code=502, detail="The generated test was invalid. Please try again.") from exc
    validate_generated_test(payload, expected_count)
    return payload


def validate_generated_test(payload: GeneratedTest, expected_count: int, requested_topic: Optional[str] = None) -> None:
    question_ids = [item.question_id for item in payload.questions]
    prompts = [item.prompt.casefold() for item in payload.questions]
    if len(payload.questions) != expected_count or len(set(question_ids)) != len(question_ids) or len(set(prompts)) != len(prompts):
        raise HTTPException(status_code=502, detail="The generated test did not contain distinct questions. Please try again.")
    if any(item.question_type != "mcq" for item in payload.questions):
        raise HTTPException(status_code=502, detail="The generated test contained an unsupported question type. Please try again.")
    if requested_topic:
        terms = {term for term in re.findall(r"[a-z0-9]+", requested_topic.casefold()) if len(term) > 2}
        generic_topics = {"core concept", "lesson concept", "general knowledge"}
        for item in payload.questions:
            searchable = f"{item.prompt} {item.topic} {item.concept or ''}".casefold()
            if terms and not any(term in searchable for term in terms) and item.topic.casefold() not in generic_topics:
                raise HTTPException(status_code=502, detail="The generated test did not match the requested topic. Please try again.")


def test_collection():
    if mongo_client is None or mongo_status != "connected":
        return None
    return mongo_client[MONGODB_DATABASE]["learning_tests"]


async def store_test(test_document: dict) -> None:
    collection = test_collection()
    if collection is None:
        fallback_active_tests[test_document["test_id"]] = test_document
        return
    try:
        await collection.insert_one(test_document)
    except Exception as exc:
        raise HTTPException(status_code=503, detail="The test could not be saved. Please try again.") from exc


async def find_test(test_id: str, session_hash: str, user_id: Optional[str] = None) -> Optional[dict]:
    ownership = {"test_id": test_id, "session_hash": session_hash}
    if user_id:
        ownership["user_id"] = user_id
    collection = test_collection()
    if collection is None:
        document = fallback_active_tests.get(test_id)
        return document if document and document.get("session_hash") == session_hash and (not user_id or document.get("user_id") == user_id) else None
    try:
        return await collection.find_one(ownership)
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Test data is temporarily unavailable. Please try again.") from exc


async def save_test_result(test_document: dict, result: dict) -> None:
    collection = test_collection()
    if collection is None:
        if test_document.get("status") != "active":
            raise HTTPException(status_code=409, detail="This test has already been submitted.")
        test_document.update(result)
        return
    try:
        updated = await collection.update_one(
            {"test_id": test_document["test_id"], "session_hash": test_document["session_hash"], "status": "active"},
            {"$set": result, "$unset": {"expires_at": ""}},
        )
        if updated.modified_count != 1:
            raise HTTPException(status_code=409, detail="This test has already been submitted.")
    except Exception as exc:
        if isinstance(exc, HTTPException):
            raise
        raise HTTPException(status_code=503, detail="The test result could not be saved. Please try again.") from exc


def normalized_skill(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().casefold())[:120]


def profile_skills(profile: dict) -> list[str]:
    values = profile_declared_skills(profile) + profile_resume_skills(profile)
    unique: list[str] = []
    for value in values:
        if value and normalized_skill(value) not in {normalized_skill(item) for item in unique}:
            unique.append(value[:120])
    return unique


def unique_skill_values(values: list[str]) -> list[str]:
    unique: list[str] = []
    for value in values:
        if value and normalized_skill(value) not in {normalized_skill(item) for item in unique}:
            unique.append(value[:120])
    return unique


def profile_declared_skills(profile: dict) -> list[str]:
    skills = profile.get("skills") or {}
    return unique_skill_values([
        item
        for key in ("programming_languages", "technical_skills", "tools_technologies", "soft_skills")
        for item in (skills.get(key) or [])
    ])


def profile_resume_skills(profile: dict) -> list[str]:
    resume_info = profile.get("resume_information") or {}
    resume = resume_info.get("extracted_information") or {}
    return unique_skill_values((resume.get("skills") or resume_info.get("skills") or []))


def skill_level_for_percentage(percentage: Optional[float], question_count: int) -> Optional[str]:
    """Product estimate only; a short assessment is not a complete proficiency measure."""
    if percentage is None or question_count < MIN_TOPIC_QUESTIONS:
        return None
    if percentage < SKILL_LEVEL_DEVELOPING_MIN:
        return "Beginner"
    if percentage < SKILL_LEVEL_PROFICIENT_MIN:
        return "Developing"
    if percentage < SKILL_LEVEL_STRONG_MIN:
        return "Proficient"
    return "Strong"


def skill_level_rules() -> dict:
    return {
        "minimum_questions": MIN_TOPIC_QUESTIONS,
        "developing_minimum": SKILL_LEVEL_DEVELOPING_MIN,
        "proficient_minimum": SKILL_LEVEL_PROFICIENT_MIN,
        "strong_minimum": SKILL_LEVEL_STRONG_MIN,
    }


def skill_level_explanation() -> str:
    return (
        f"Product estimate only: Beginner is below {SKILL_LEVEL_DEVELOPING_MIN}%, "
        f"Developing is {SKILL_LEVEL_DEVELOPING_MIN}–{SKILL_LEVEL_PROFICIENT_MIN - 0.1:g}%, "
        f"Proficient is {SKILL_LEVEL_PROFICIENT_MIN}–{SKILL_LEVEL_STRONG_MIN - 0.1:g}%, "
        f"and Strong is {SKILL_LEVEL_STRONG_MIN}% or higher. At least {MIN_TOPIC_QUESTIONS} "
        "questions are needed; a short test is not a complete measure of ability."
    )


def topic_analysis_from_review(review: list[dict]) -> list[dict]:
    grouped: dict[str, dict] = {}
    for item in review:
        topic = item["topic"]
        stats = grouped.setdefault(topic, {"topic": topic, "total_questions": 0, "correct_count": 0, "incorrect_count": 0, "unanswered_count": 0})
        stats["total_questions"] += 1
        if item["is_correct"]:
            stats["correct_count"] += 1
        elif item.get("selected_answer") is None:
            stats["unanswered_count"] += 1
        else:
            stats["incorrect_count"] += 1
    analysis = []
    for stats in grouped.values():
        percentage = round(stats["correct_count"] / stats["total_questions"] * 100, 1)
        level = skill_level_for_percentage(percentage, stats["total_questions"])
        analysis.append({**stats, "percentage": percentage, "level": level, "status": "assessed" if level else "not_enough_data"})
    return analysis


def enrich_result_analysis(result: dict, topic_analysis: list[dict], review: list[dict]) -> dict:
    assessed = [item for item in topic_analysis if item["level"]]
    strong_topics = [item["topic"] for item in assessed if item["level"] == "Strong"]
    weak_topics = [item["topic"] for item in topic_analysis if item["percentage"] < PASSING_SCORE_PERCENTAGE]
    return {
        **result,
        "topic_analysis": topic_analysis,
        "strong_topics": strong_topics,
        "weak_topics": weak_topics,
        "concepts_to_review": unique_skill_values([
            (item.get("concept") or item["topic"])
            for item in review
            if not item["is_correct"]
        ]),
        "skill_level_rules": skill_level_rules(),
        "estimate_note": skill_level_explanation(),
    }


async def save_assessment_attempt(test_document: dict, result: dict, user_id: str) -> None:
    attempt = {"attempt_id": f"attempt_{uuid.uuid4().hex[:16]}", "assessment_id": test_document["test_id"], "test_id": test_document["test_id"], "user_id": user_id, "topic": test_document["topic"], "subject": test_document.get("subject"), "difficulty": test_document["difficulty"], "total_questions": result["total_questions"], "correct_count": result["correct_count"], "incorrect_count": result["incorrect_count"], "unanswered_count": result["unanswered_count"], "score": result["score"], "percentage": result["percentage"], "passed": result["passed"], "weak_topics": result["weak_topics"], "strong_topics": result["strong_topics"], "topic_analysis": result["topic_analysis"], "submitted_at": result["submitted_at"]}
    collection = test_collection()
    if collection is None:
        return
    try:
        await mongo_client[MONGODB_DATABASE]["assessment_attempts"].insert_one(attempt)
    except Exception as exc:
        raise HTTPException(status_code=503, detail="The assessment result could not be archived. Please try again.") from exc


async def save_skill_verification(test_document: dict, result: dict, user_id: str, profile: dict) -> Optional[dict]:
    if not result["passed"]:
        return None
    skill = test_document.get("subject") or test_document["topic"]
    verification = {"user_id": user_id, "skill": skill, "skill_key": normalized_skill(skill), "assessment_id": test_document["test_id"], "score": result["percentage"], "level": test_document["difficulty"], "verified_at": result["submitted_at"]}
    collection = test_collection()
    if collection is None:
        current = fallback_skill_verifications.setdefault(user_id, [])
        current[:] = [item for item in current if item["skill_key"] != verification["skill_key"]]
        current.append(verification)
        return verification
    try:
        await mongo_client[MONGODB_DATABASE]["skill_verifications"].update_one(
            {"user_id": user_id, "skill_key": verification["skill_key"]}, {"$set": verification}, upsert=True
        )
        return verification
    except Exception as exc:
        raise HTTPException(status_code=503, detail="The verified skill could not be saved. Please try again.") from exc


def build_learning_plan(test_document: dict, result: dict, profile: dict, review: list[dict]) -> dict:
    weak_topics = result["weak_topics"]
    tasks = []
    for topic in weak_topics:
        explanation = next((item["explanation"] for item in review if item["topic"] == topic), "Review the core ideas and practice a new example.")
        tasks.append({"topic": topic, "task": f"Review {topic} and solve one new example.", "estimated_minutes": 30, "guidance": explanation, "completed": False})
    skills = profile_skills(profile)
    current = normalized_skill(test_document.get("subject") or test_document["topic"])
    next_topic = next((skill for skill in skills if normalized_skill(skill) != current), None)
    if result["passed"]:
        recommendation = {"action": "Take a new assessment", "topic": next_topic or test_document["topic"], "difficulty": "Advanced" if test_document["difficulty"] in {"Hard", "Advanced"} else "Hard"}
    else:
        recommendation = {"action": "Review the recommended tasks, then retest", "topic": weak_topics[0] if weak_topics else test_document["topic"], "difficulty": test_document["difficulty"]}
    return {"plan_id": f"plan_{uuid.uuid4().hex[:16]}", "assessment_id": test_document["test_id"], "user_id": profile["user_id"], "status": "active", "created_at": result["submitted_at"], "passed": result["passed"], "tasks": tasks, "recommended_next_assessment": recommendation}


async def save_learning_plan(plan: dict) -> dict:
    collection = test_collection()
    if collection is None:
        fallback_learning_plans.setdefault(plan["user_id"], []).append(plan)
        return plan
    try:
        await mongo_client[MONGODB_DATABASE]["learning_plans"].insert_one(plan)
        return plan
    except Exception as exc:
        raise HTTPException(status_code=503, detail="The learning plan could not be saved. Please try again.") from exc


ROLE_SKILLS = {
    "software developer": ["Programming fundamentals", "Python", "C++", "Data structures and algorithms", "Object-oriented programming", "DBMS and SQL", "APIs and backend development", "Git and GitHub", "HTML, CSS, and JavaScript"],
    "software engineer": ["Programming fundamentals", "Python", "C++", "Data structures and algorithms", "Object-oriented programming", "DBMS and SQL", "APIs and backend development", "Git and GitHub", "HTML, CSS, and JavaScript"],
    "python developer": ["Python", "SQL", "Git", "APIs"],
    "data analyst": ["SQL", "Python", "Statistics", "Data visualization"],
    "frontend developer": ["HTML", "CSS", "JavaScript", "Web development"],
    "full stack developer": ["HTML", "CSS", "JavaScript", "SQL", "APIs"],
}

SKILL_PATHS = {
    "programming fundamentals": [
        ("Variables and data types", ["Identify data types in small examples", "Practice assigning and converting values"]),
        ("Conditions and loops", ["Trace conditional branches", "Practice loops with a stopping condition"]),
        ("Functions", ["Write a function with inputs and a return value", "Practice calling and testing functions"]),
    ],
    "python": [
        ("Variables and data types", ["Identify Python data types in examples", "Practice assignment and type conversion"]),
        ("Conditions and loops", ["Trace if/elif/else branches", "Practice for and while loops"]),
        ("Functions", ["Write functions with parameters and return values", "Practice scope and reusable helpers"]),
        ("Lists, tuples, sets, and dictionaries", ["Choose a collection for a given use", "Practice common collection operations"]),
        ("Object-oriented programming", ["Identify classes, objects, and methods", "Practice encapsulation and inheritance"]),
        ("Exception handling", ["Identify common exceptions", "Practice try/except and meaningful recovery"]),
        ("File handling", ["Practice reading and writing files safely", "Use context managers for file operations"]),
    ],
    "c++": [
        ("Variables and data types", ["Practice declarations and type conversions", "Trace values through a short program"]),
        ("Conditions and loops", ["Trace conditional branches", "Practice loops and termination"]),
        ("Functions", ["Practice parameters, returns, and overloads", "Write a small reusable function"]),
        ("Collections and strings", ["Practice arrays, vectors, and strings", "Select a collection for a small task"]),
        ("Object-oriented programming", ["Practice classes and access control", "Review constructors and inheritance"]),
        ("Pointers and memory", ["Trace pointers and references", "Practice safe resource ownership"]),
        ("Exception handling", ["Review exception flow", "Practice handling expected failures"]),
    ],
    "data structures and algorithms": [
        ("Arrays and strings", ["Practice indexing and traversal", "Solve a simple array transformation"]),
        ("Linked lists and stacks", ["Trace insertions and removals", "Practice stack operations"]),
        ("Queues, trees, and graphs", ["Compare core structures", "Trace a traversal on a small example"]),
        ("Searching and sorting", ["Compare search strategies", "Trace a sorting algorithm"]),
        ("Complexity analysis", ["Count operations in a loop", "Compare time and space complexity"]),
    ],
    "object-oriented programming": [
        ("Classes and objects", ["Identify class and instance responsibilities", "Design a small class"]),
        ("Encapsulation and abstraction", ["Practice hiding implementation details", "Use clear public interfaces"]),
        ("Inheritance and composition", ["Compare inheritance and composition", "Model a relationship with a small example"]),
        ("Polymorphism", ["Trace method dispatch", "Practice substitutable implementations"]),
    ],
    "dbms and sql": [
        ("Relational data and keys", ["Identify primary and foreign keys", "Model a small relational schema"]),
        ("SQL queries and filtering", ["Practice SELECT and WHERE", "Use ordering and aggregation"]),
        ("Joins and aggregation", ["Choose joins for a question", "Practice GROUP BY and aggregate queries"]),
        ("Normalization and transactions", ["Recognize basic normalization", "Review transaction guarantees"]),
        ("Indexes and query plans", ["Explain what an index supports", "Compare simple query plans"]),
    ],
    "apis and backend development": [
        ("HTTP and REST fundamentals", ["Identify HTTP methods and status codes", "Design a resource-oriented endpoint"]),
        ("Request validation and errors", ["Validate a structured request", "Return actionable error responses"]),
        ("Authentication and authorization", ["Distinguish identity from access control", "Apply ownership checks to an API"]),
        ("Persistence and API testing", ["Connect an endpoint to saved data", "Test success and failure cases"]),
    ],
    "git and github": [
        ("Repositories and commits", ["Practice staging and committing a focused change", "Inspect commit history"]),
        ("Branches and merges", ["Create and switch branches", "Resolve a small merge conflict"]),
        ("Pull requests and review", ["Open a pull request", "Respond to review feedback"]),
        ("Remote collaboration", ["Sync with a remote safely", "Review differences before pushing"]),
    ],
    "html, css, and javascript": [
        ("Semantic HTML", ["Build a semantic page structure", "Label inputs accessibly"]),
        ("CSS layout and responsive design", ["Practice flex and grid layouts", "Test a layout at mobile widths"]),
        ("JavaScript fundamentals", ["Practice functions and DOM events", "Validate user input and display state"]),
        ("Browser APIs and accessibility", ["Use browser APIs for an interaction", "Check keyboard and screen-reader behavior"]),
    ],
}

SKILL_MAP_ROLES = {
    "Software Developer": ROLE_SKILLS["software developer"],
    "Software Engineer": ROLE_SKILLS["software engineer"],
    "Python Developer": ROLE_SKILLS["python developer"],
    "Data Analyst": ROLE_SKILLS["data analyst"],
    "Frontend Developer": ROLE_SKILLS["frontend developer"],
    "Full Stack Developer": ROLE_SKILLS["full stack developer"],
}

SKILL_PATH_ALIASES = {
    "sql": "dbms and sql",
    "apis": "apis and backend development",
    "git": "git and github",
    "web development": "html, css, and javascript",
}


def skill_topic_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", normalized_skill(value)).strip("-")


def target_role_and_requirements(profile: dict) -> tuple[Optional[str], list[str]]:
    career = profile.get("career_information") or {}
    goals = profile.get("career_goals") or {}
    role = (goals.get("target_role") or career.get("target_job_role") or "").strip() or None
    if not role:
        return None, []
    role_key = role.casefold()
    required = next((skills for name, skills in ROLE_SKILLS.items() if name in role_key), [])
    return role, required


def profile_learning_goals(profile: dict) -> list[str]:
    goals = profile.get("career_goals") or {}
    return unique_skill_values([
        goals.get("short_term_career_goal") or "",
        goals.get("long_term_career_goal") or "",
        goals.get("areas_to_improve") or "",
    ])


def assessment_topic_result(attempts: list[dict], topic_name: str) -> Optional[dict]:
    topic_key = normalized_skill(topic_name)
    for attempt in sorted(
        attempts,
        key=lambda item: item.get("submitted_at") or datetime.min.replace(tzinfo=timezone.utc),
        reverse=True,
    ):
        for result in attempt.get("topic_analysis", []):
            result_key = normalized_skill(result.get("topic", ""))
            if result_key == topic_key:
                return {
                    "assessment_id": attempt.get("assessment_id") or attempt.get("test_id"),
                    "percentage": result.get("percentage"),
                    "total_questions": result.get("total_questions", 0),
                    "passed": (
                        int(result.get("total_questions", 0)) >= MIN_TOPIC_QUESTIONS
                        and float(result.get("percentage", 0)) >= PASSING_SCORE_PERCENTAGE
                    ),
                    "submitted_at": attempt.get("submitted_at"),
                }
    return None


def skill_assessment_result(attempts: list[dict], skill: str) -> Optional[dict]:
    skill_key = normalized_skill(skill)
    matches = [
        attempt for attempt in attempts
        if normalized_skill(attempt.get("subject") or "") == skill_key
        or normalized_skill(attempt.get("topic") or "") == skill_key
    ]
    if not matches:
        return None
    latest = max(
        matches,
        key=lambda item: item.get("submitted_at") or datetime.min.replace(tzinfo=timezone.utc),
    )
    return {
        "assessment_id": latest.get("assessment_id") or latest.get("test_id"),
        "score": latest.get("percentage"),
        "passed": bool(latest.get("passed")),
        "submitted_at": latest.get("submitted_at"),
    }


async def load_topic_progress(user_id: str) -> list[dict]:
    collection = test_collection()
    if collection is None:
        return list(fallback_topic_learning_progress.get(user_id, {}).values())
    try:
        cursor = mongo_client[MONGODB_DATABASE]["topic_learning_progress"].find(
            {"user_id": user_id},
            {"_id": 0},
        )
        return [item async for item in cursor]
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Learning topic progress is temporarily unavailable.") from exc


async def save_topic_progress_record(user_id: str, record: dict) -> None:
    collection = test_collection()
    if collection is None:
        fallback_topic_learning_progress.setdefault(user_id, {})[
            (record["skill_key"], record["topic_key"])
        ] = record
        return
    try:
        await mongo_client[MONGODB_DATABASE]["topic_learning_progress"].update_one(
            {
                "user_id": user_id,
                "skill_key": record["skill_key"],
                "topic_key": record["topic_key"],
            },
            {"$set": record},
            upsert=True,
        )
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Learning topic progress could not be saved.") from exc


async def refresh_topic_progress_status(user_id: str, attempts: list[dict]) -> None:
    for record in await load_topic_progress(user_id):
        assessment = assessment_topic_result(attempts, record["topic"])
        assessment_status = "not_assessed"
        if assessment:
            assessment_status = "passed" if assessment["passed"] else "needs_improvement"
        all_tasks_complete = bool(record.get("tasks")) and all(
            task.get("completed") for task in record["tasks"]
        )
        completed = all_tasks_complete and assessment_status == "passed"
        any_task_complete = any(task.get("completed") for task in record.get("tasks", []))
        status = (
            "completed" if completed
            else "in_progress" if any_task_complete or assessment
            else "not_started"
        )
        updated = {
            **record,
            "assessment_status": assessment_status,
            "assessment_score": assessment.get("percentage") if assessment else None,
            "assessment_id": assessment.get("assessment_id") if assessment else None,
            "status": status,
            "completed_at": record.get("completed_at") or (datetime.now(timezone.utc) if completed else None),
            "updated_at": datetime.now(timezone.utc),
        }
        if not completed:
            updated["completed_at"] = None
        await save_topic_progress_record(user_id, updated)


async def skill_path_data(profile: dict, selected_role: Optional[str] = None) -> dict:
    user_id = profile["user_id"]
    configured_role = selected_role or target_role_and_requirements(profile)[0]
    role = configured_role or "Software Developer"
    role_name = next((name for name in SKILL_MAP_ROLES if name.casefold() == role.casefold()), None)
    if role_name is None:
        if selected_role:
            raise HTTPException(status_code=422, detail="Choose a supported target role to view its skill map.")
        role_name = "Software Developer"
    role_is_fallback = configured_role is not None and role_name == "Software Developer" and configured_role.casefold() != role_name.casefold()
    skill_names = SKILL_MAP_ROLES[role_name]
    attempts, verified, _ = await dashboard_data(profile)
    progress_records = await load_topic_progress(user_id)
    progress_by_key = {
        (item.get("skill_key"), item.get("topic_key")): item
        for item in progress_records
    }
    skill_progress = []
    qualified = []
    for skill in skill_names:
        skill_key = normalized_skill(skill)
        configured_topics = SKILL_PATHS.get(SKILL_PATH_ALIASES.get(skill_key, skill_key))
        if configured_topics is None:
            # Existing role maps may include simple skills without a multi-topic path.
            configured_topics = [(skill, [f"Complete a focused learning activity for {skill}."])]
        topics = []
        for topic_name, task_names in configured_topics:
            record = progress_by_key.get((normalized_skill(skill), skill_topic_key(topic_name)), {})
            saved_tasks = {item.get("task_id"): item for item in record.get("tasks", [])}
            tasks = [
                {
                    "task_id": f"task-{index + 1}",
                    "description": task_name,
                    "completed": bool(saved_tasks.get(f"task-{index + 1}", {}).get("completed")),
                }
                for index, task_name in enumerate(task_names)
            ]
            assessment = assessment_topic_result(attempts, topic_name)
            assessment_status = "not_assessed"
            if assessment:
                assessment_status = "passed" if assessment["passed"] else "needs_improvement"
            all_tasks_complete = bool(tasks) and all(item["completed"] for item in tasks)
            completed = all_tasks_complete and assessment_status == "passed"
            any_task_complete = any(item["completed"] for item in tasks)
            status = (
                "completed" if completed
                else "in_progress" if any_task_complete or assessment
                else "not_started"
            )
            topics.append({
                "topic": topic_name,
                "status": status,
                "tasks": tasks,
                "assessment_status": assessment_status,
                "assessment_score": assessment["percentage"] if assessment else None,
                "assessment_id": assessment["assessment_id"] if assessment else None,
                "completed_at": serialize_document(record.get("completed_at")) if completed else None,
            })
        required_topics_complete = bool(topics) and all(item["status"] == "completed" for item in topics)
        skill_assessment = skill_assessment_result(attempts, skill)
        skill_verified = any(normalized_skill(item.get("skill", "")) == normalized_skill(skill) for item in verified)
        eligible = bool(required_topics_complete and skill_assessment and skill_assessment["passed"])
        if eligible:
            qualified.append({
                "skill": skill,
                "assessment_id": skill_assessment["assessment_id"],
                "score": skill_assessment["score"],
                "completed_at": max(
                    (item["completed_at"] for item in topics if item["completed_at"]),
                    default=serialize_document(skill_assessment["submitted_at"]),
                ),
            })
        skill_progress.append({
            "skill": skill,
            "status": "verified" if skill_verified else "passed_assessment" if skill_assessment and skill_assessment["passed"] else "needs_improvement" if skill_assessment else "in_progress" if any(item["status"] != "not_started" for item in topics) else "not_assessed",
            "topics": topics,
            "required_topics_complete": required_topics_complete,
            "assessment_status": "not_assessed" if not skill_assessment else "passed" if skill_assessment["passed"] else "needs_improvement",
            "assessment_score": skill_assessment["score"] if skill_assessment else None,
            "assessment_id": skill_assessment["assessment_id"] if skill_assessment else None,
            "eligible_for_resume": eligible,
            "already_in_resume": normalized_skill(skill) in {
                normalized_skill(item) for item in profile_resume_skills(profile)
            },
            "verified": skill_verified,
        })
    return {
        "target_role": role_name,
        "requested_target_role": configured_role,
        "available_roles": list(SKILL_MAP_ROLES),
        "skills": skill_progress,
        "qualified_skills": qualified,
        "passing_percentage": PASSING_SCORE_PERCENTAGE,
        "message": (
            f"No configured skill map exists for {configured_role}; showing the Software Developer map."
            if role_is_fallback
            else "Complete each topic's learning tasks and pass its topic assessment; then pass the skill assessment. Resume inclusion always requires your confirmation."
        ),
    }


def build_dashboard(profile: dict, attempts: list[dict], verified: list[dict], plan: Optional[dict]) -> dict:
    attempts = sorted(attempts, key=lambda item: item.get("submitted_at") or datetime.min.replace(tzinfo=timezone.utc), reverse=True)
    topic_totals: dict[str, dict] = {}
    for attempt in attempts:
        for topic in attempt.get("topic_analysis", []):
            current = topic_totals.setdefault(topic["topic"], {"topic": topic["topic"], "total_questions": 0, "correct_count": 0, "incorrect_count": 0, "unanswered_count": 0, "assessment_count": 0})
            current["assessment_count"] += 1
            for field in ("total_questions", "correct_count", "incorrect_count", "unanswered_count"):
                current[field] += int(topic.get(field, 0))
    topic_summary = []
    for item in topic_totals.values():
        percentage = round(item["correct_count"] / item["total_questions"] * 100, 1) if item["total_questions"] else None
        level = skill_level_for_percentage(percentage, item["total_questions"])
        topic_summary.append({**item, "percentage": percentage, "level": level, "status": "assessed" if level else "not_enough_data"})
    assessed = [item for item in topic_summary if item["level"]]
    strongest = max(assessed, key=lambda item: item["percentage"]) if assessed else None
    weakest = min(assessed, key=lambda item: item["percentage"]) if assessed else None
    role, required = target_role_and_requirements(profile)
    assessed_names = {normalized_skill(item["topic"]) for item in topic_summary}
    verified_names = {normalized_skill(item.get("skill", "")) for item in verified}
    not_assessed = [skill for skill in required if normalized_skill(skill) not in assessed_names and normalized_skill(skill) not in verified_names]
    skills_to_improve = [item for item in topic_summary if item["percentage"] < PASSING_SCORE_PERCENTAGE]
    required_needing_improvement = [
        item["topic"]
        for item in skills_to_improve
        if normalized_skill(item["topic"]) in {normalized_skill(skill) for skill in required}
    ]
    history = [{key: serialize_document(item.get(key)) for key in ("test_id", "topic", "difficulty", "score", "percentage", "passed", "submitted_at", "weak_topics")} for item in attempts[:20]]
    passed = sum(1 for item in attempts if item.get("passed"))
    latest = attempts[0] if attempts else None
    scored_attempts = [item for item in attempts if item.get("percentage") is not None]
    learning_plan_progress = None
    if plan:
        tasks = plan.get("tasks", [])
        completed_tasks = sum(1 for item in tasks if item.get("completed"))
        learning_plan_progress = {
            "completed_tasks": completed_tasks,
            "total_tasks": len(tasks),
            "percentage": round(completed_tasks / len(tasks) * 100) if tasks else 0,
        }
    if weakest or skills_to_improve:
        topic_to_review = weakest or min(skills_to_improve, key=lambda item: item["percentage"])
        next_action = f"Review {topic_to_review['topic']} and retake an assessment."
    elif not attempts:
        next_action = "Take your first assessment to establish a baseline."
    elif not_assessed:
        next_action = f"Assess {not_assessed[0]} for your {role or 'target'} goal."
    elif not role:
        next_action = "Choose a target role or continue with a general skill assessment."
    else:
        next_action = "Continue your learning plan or take a more challenging assessment."
    return {
        "overview": {"latest_score": latest.get("score") if latest else None, "latest_percentage": latest.get("percentage") if latest else None, "average_percentage": round(sum(float(item["percentage"]) for item in scored_attempts) / len(scored_attempts), 1) if scored_attempts else None, "completed_count": len(attempts), "passed_count": passed, "failed_count": len(attempts) - passed, "passing_percentage": PASSING_SCORE_PERCENTAGE},
        "topic_summary": sorted(topic_summary, key=lambda item: item["topic"].casefold()),
        "strongest_topic": strongest,
        "weakest_topic": weakest,
        "skills_to_improve": skills_to_improve,
        "verified_skills": [{key: serialize_document(item.get(key)) for key in ("skill", "score", "level", "verified_at", "assessment_id")} for item in verified],
        "profile_skills": profile_declared_skills(profile),
        "resume_skills": profile_resume_skills(profile),
        "skill_gap": {
            "target_role": role,
            "required_skills": required,
            "not_assessed": not_assessed,
            "assessed_needing_improvement": required_needing_improvement,
            "learning_goals": profile_learning_goals(profile),
            "message": (
                "Select a target role in your profile to compare required skills."
                if not role
                else "No role-specific skill map is available yet; use the general topic analysis."
                if not required
                else None
            ),
        },
        "learning_plan": serialize_document(plan) if plan else None,
        "learning_plan_progress": learning_plan_progress,
        "history": history,
        "score_trend": [
            {
                "topic": item.get("topic", "Assessment"),
                "percentage": item.get("percentage"),
                "submitted_at": serialize_document(item.get("submitted_at")),
            }
            for item in sorted(
                attempts,
                key=lambda item: item.get("submitted_at") or datetime.min.replace(tzinfo=timezone.utc),
            )[-20:]
        ],
        "next_action": next_action,
        "skill_level_rules": skill_level_rules(),
        "estimate_note": skill_level_explanation(),
    }


async def dashboard_data(profile: dict) -> tuple[list[dict], list[dict], Optional[dict]]:
    collection = test_collection()
    user_id = profile["user_id"]
    if collection is None:
        attempts = [item for item in fallback_active_tests.values() if item.get("user_id") == user_id and item.get("status") == "submitted"]
        verified = fallback_skill_verifications.get(user_id, [])
        plans = fallback_learning_plans.get(user_id, [])
        return attempts, verified, plans[-1] if plans else None
    try:
        attempts = [item async for item in mongo_client[MONGODB_DATABASE]["assessment_attempts"].find({"user_id": user_id}, {"_id": 0}).sort("submitted_at", -1)]
        verified = [item async for item in mongo_client[MONGODB_DATABASE]["skill_verifications"].find({"user_id": user_id}, {"_id": 0, "user_id": 0}).sort("verified_at", -1)]
        plan = await mongo_client[MONGODB_DATABASE]["learning_plans"].find_one({"user_id": user_id}, {"_id": 0}, sort=[("created_at", -1)])
        return attempts, verified, plan
    except Exception as exc:
        raise HTTPException(status_code=503, detail="The student dashboard is temporarily unavailable. Please try again.") from exc


@app.post("/api/tests/generate")
async def generate_test(request: TestGenerateRequest, http_request: Request):
    use_sample_quiz = request.use_mock and ENABLE_MOCK_QUIZ
    if gemini_client is None and not use_sample_quiz:
        raise HTTPException(status_code=503, detail="Gemini is not configured. Set GEMINI_API_KEY in the backend environment.")
    profile = await require_profile(http_request)
    allowed, _ = await check_rate_limit(http_request)
    if not allowed:
        raise HTTPException(status_code=429, detail="You have reached the chat request limit. Please try again later.", headers={"Retry-After": str(CHAT_RATE_LIMIT_WINDOW)})
    conversation = "\n".join(f"{item.role}: {item.content}" for item in request.history)
    if not request.topic and not request.subject and len(conversation.strip()) < 80:
        raise HTTPException(status_code=422, detail="Tell me the lesson topic or add more lesson conversation before starting a test.")
    topic_context = f"Primary subject/topic to teach and assess: {request.subject or request.topic}" if (request.subject or request.topic) else ""
    student_context = profile_context(profile)
    conversation_context = f"Supporting lesson conversation:\n{conversation}" if conversation.strip() else ""
    lesson_context = "\n\n".join(part for part in (topic_context, student_context, conversation_context) if part)[:TEST_MAX_LESSON_CHARS]
    if use_sample_quiz:
        payload = mock_test_payload(request)
        source = "mock_sample"
    else:
        try:
            payload = await generate_test_payload(test_prompt(request, lesson_context), request.question_count)
            validate_generated_test(payload, request.question_count, request.topic or request.subject)
            source = "gemini"
        except HTTPException as exc:
            if ENABLE_MOCK_QUIZ and exc.status_code in {429, 502, 503, 504}:
                logger.warning("Using local quiz fallback after Gemini failure: status=%s", exc.status_code)
                payload = mock_test_payload(request)
                source = "mock_fallback"
            else:
                raise
    stored_questions = [test_question_for_storage(question) for question in payload.questions]
    for question in stored_questions:
        question["difficulty"] = request.difficulty
        question["question_type"] = request.question_type
    now = datetime.now(timezone.utc)
    test_document = {
        "test_id": f"test_{uuid.uuid4().hex[:16]}",
        "session_hash": session_key(http_request),
        "topic": request.topic or request.subject or payload.questions[0].topic,
        "subject": request.subject or request.topic,
        "difficulty": request.difficulty,
        "question_count": request.question_count,
        "questions": stored_questions,
        "user_id": profile["user_id"],
        "source": source,
        "status": "active",
        "created_at": now,
        "expires_at": now + timedelta(seconds=TEST_ACTIVE_TTL_SECONDS),
    }
    await store_test(test_document)
    return public_test(test_document)


@app.post("/api/tests/submit")
async def submit_test(request: TestSubmitRequest, http_request: Request):
    profile = await require_profile(http_request)
    test_document = await find_test(request.test_id, session_key(http_request), profile["user_id"])
    if test_document is None:
        raise HTTPException(status_code=404, detail="Test not found or it belongs to another session.")
    if test_document.get("status") != "active":
        raise HTTPException(status_code=409, detail="This test has already been submitted.")
    expires_at = test_document.get("expires_at")
    if expires_at is not None:
        if expires_at.tzinfo is None:
            expires_at = expires_at.replace(tzinfo=timezone.utc)
        if datetime.now(timezone.utc) >= expires_at:
            raise HTTPException(status_code=410, detail="This test has expired. Please generate a new test.")
    questions = test_document["questions"]
    valid_ids = {item["question_id"] for item in questions}
    if set(request.answers) - valid_ids:
        raise HTTPException(status_code=422, detail="The submitted answers do not match this test.")
    correct_count = 0
    review: list[dict] = []
    for question in questions:
        selected = request.answers.get(question["question_id"])
        is_correct = selected == question["correct_answer"]
        if is_correct:
            correct_count += 1
        review.append({**question, "selected_answer": selected, "is_correct": is_correct})
    submitted_at = datetime.now(timezone.utc)
    created_at = test_document["created_at"]
    if created_at.tzinfo is None:
        created_at = created_at.replace(tzinfo=timezone.utc)
    time_taken = max(0, int((submitted_at - created_at).total_seconds()))
    total = len(questions)
    percentage = round(correct_count / total * 100, 1)
    topic_analysis = topic_analysis_from_review(review)
    incorrect_count = sum(item["incorrect_count"] for item in topic_analysis)
    unanswered_count = sum(item["unanswered_count"] for item in topic_analysis)
    result = {
        "status": "submitted",
        "total_questions": total,
        "correct_count": correct_count,
        "incorrect_count": incorrect_count,
        "unanswered_count": unanswered_count,
        "score": f"{correct_count}/{total}",
        "percentage": percentage,
        "passed": percentage >= PASSING_SCORE_PERCENTAGE,
        "passing_percentage": PASSING_SCORE_PERCENTAGE,
        "time_taken_seconds": time_taken,
        "submitted_at": submitted_at,
    }
    result = enrich_result_analysis(result, topic_analysis, review)
    await save_test_result(test_document, result)
    await save_assessment_attempt(test_document, result, profile["user_id"])
    attempts, _, _ = await dashboard_data(profile)
    await refresh_topic_progress_status(profile["user_id"], attempts)
    verification = await save_skill_verification(test_document, result, profile["user_id"], profile)
    plan = await save_learning_plan(build_learning_plan(test_document, result, profile, review))
    return {"test_id": request.test_id, **result, "verified_skill": verification, "learning_plan": serialize_document(plan), "review": review}


@app.get("/api/tests/history")
async def test_history(http_request: Request):
    profile = await require_profile(http_request)
    collection = test_collection()
    session_hash = session_key(http_request)
    user_id = profile["user_id"]
    ownership = {"session_hash": session_hash, "status": "submitted"}
    if user_id:
        ownership["user_id"] = user_id
    if collection is None:
        documents = [item for item in fallback_active_tests.values() if item.get("session_hash") == session_hash and item.get("status") == "submitted" and (not user_id or item.get("user_id") == user_id)]
        documents.sort(key=lambda item: item.get("created_at", datetime.min.replace(tzinfo=timezone.utc)), reverse=True)
        return {"tests": [serialize_document({key: item.get(key) for key in ("test_id", "topic", "difficulty", "question_count", "score", "percentage", "passed", "weak_topics", "created_at", "submitted_at")}) for item in documents[:50]]}
    try:
        cursor = collection.find(ownership, {"_id": 0, "session_hash": 0, "questions": 0, "user_id": 0}).sort("submitted_at", -1).limit(50)
        documents = [serialize_document(document) async for document in cursor]
        return {"tests": documents}
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Test history is temporarily unavailable. Please try again.") from exc


@app.get("/api/tests/skills")
async def verified_skills(http_request: Request):
    profile = await require_profile(http_request)
    user_id = profile["user_id"]
    collection = test_collection()
    if collection is None:
        return {"skills": [serialize_document(item) for item in fallback_skill_verifications.get(user_id, [])]}
    try:
        cursor = mongo_client[MONGODB_DATABASE]["skill_verifications"].find({"user_id": user_id}, {"_id": 0, "user_id": 0}).sort("verified_at", -1)
        return {"skills": [serialize_document(item) async for item in cursor]}
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Verified skills are temporarily unavailable. Please try again.") from exc


@app.get("/api/tests/dashboard")
async def assessment_dashboard(http_request: Request):
    profile = await require_profile(http_request)
    attempts, verified, plan = await dashboard_data(profile)
    dashboard = build_dashboard(profile, attempts, verified, plan)
    dashboard["skill_map"] = await skill_path_data(profile)
    return dashboard


def valid_adzuna_url(value: object) -> bool:
    if not isinstance(value, str):
        return False
    parsed = urlparse(value)
    host = (parsed.hostname or "").lower()
    allowed_roots = ("adzuna.com", "adzuna.co.uk", "adzuna.co.in", "adzuna.com.au", "adzuna.com.br", "adzuna.fr", "adzuna.de", "adzuna.nl", "adzuna.it", "adzuna.es", "adzuna.co.za", "adzuna.sg", "adzuna.us")
    allowed_host = any(host == root or host.endswith(f".{root}") for root in allowed_roots)
    try:
        safe_port = parsed.port in {None, 80, 443}
    except ValueError:
        return False
    return parsed.scheme in {"http", "https"} and allowed_host and safe_port and parsed.username is None and parsed.password is None


def secure_adzuna_url(value: str) -> str:
    parsed = urlparse(value)
    return parsed._replace(scheme="https", netloc=parsed.netloc.lower()).geturl()


def fetch_adzuna_jobs(
    role: str,
    location: str,
    country: str,
    remote: Optional[bool] = None,
    prioritize_entry_level: bool = False,
    skill_terms: Optional[list[str]] = None,
) -> list[dict]:
    search_terms = " ".join([role, *(skill_terms or [])])
    if prioritize_entry_level:
        search_terms = f"{search_terms} entry level fresher"
    if remote:
        search_terms = f"{search_terms} remote"
    query = urlencode({
        "app_id": ADZUNA_APP_ID,
        "app_key": ADZUNA_APP_KEY,
        "what": search_terms,
        **({"where": location} if location else {}),
        "results_per_page": 10,
        "max_days_old": 14,
        "sort_by": "date",
        "content-type": "application/json",
    })
    url = f"https://api.adzuna.com/v1/api/jobs/{country}/search/1?{query}"
    request = UrlRequest(url, headers={"Accept": "application/json", "User-Agent": "VaishusAIStudentAssistant/1.0"})
    try:
        with urlopen(request, timeout=12) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, OSError, ValueError) as exc:
        logger.warning("Adzuna job search failed: %s", type(exc).__name__)
        raise HTTPException(status_code=502, detail="The configured job provider could not be reached. Please try again later.") from exc
    jobs = []
    seen = set()
    for item in payload.get("results", []):
        raw_apply_url = item.get("redirect_url")
        if not valid_adzuna_url(raw_apply_url):
            continue
        apply_url = secure_adzuna_url(raw_apply_url)
        source_id = str(item.get("id") or "").strip()
        dedupe_key = source_id or apply_url.rstrip("/").casefold()
        if dedupe_key in seen:
            continue
        seen.add(dedupe_key)
        company = item.get("company") or {}
        location_data = item.get("location") or {}
        jobs.append({
            "source_id": source_id or None,
            "title": str(item.get("title") or "Title not supplied")[:250],
            "company": str(company.get("display_name") or "Company not supplied")[:200],
            "location": str(location_data.get("display_name") or "Location not supplied")[:200],
            "description": str(item.get("description") or "")[:4000],
            "experience_level": "Not supplied by provider",
            "required_skills": None,
            "posted_at": item.get("created"),
            "source": "Adzuna",
            "source_url": "https://www.adzuna.com/",
            "application_url": apply_url,
        })
        if len(jobs) == 10:
            break
    return jobs


async def daily_job_recommendations(profile: dict) -> dict:
    role_default = (
        profile.get("career_goals", {}).get("target_role")
        or profile.get("career_information", {}).get("target_job_role")
        or ""
    )
    location_default = profile.get("career_information", {}).get("preferred_work_location") or ""
    user_id = profile["user_id"]
    today = datetime.now(timezone.utc).date().isoformat()
    if mongo_client is None or mongo_status != "connected":
        raise HTTPException(status_code=503, detail="Job recommendations require the connected student database.")
    db = mongo_client[MONGODB_DATABASE]
    try:
        preferences = await db["job_preferences"].find_one({"user_id": user_id}, {"_id": 0}) or {}
        role = str(preferences.get("target_role") or role_default).strip()
        location = str(preferences.get("location") or location_default).strip()
        career = profile.get("career_information", {})
        current_status = str(career.get("current_status") or "").casefold()
        experience_level = str(career.get("experience_level") or "").casefold()
        prioritize_entry_level = any(term in current_status or term in experience_level for term in ("student", "fresher", "entry"))
        cached = await db["job_recommendations"].find_one({"user_id": user_id, "day_key": today}, {"_id": 0})
        configured = bool(ADZUNA_ENABLED and ADZUNA_APP_ID and ADZUNA_APP_KEY)
        if cached:
            return {
                "configured": configured,
                "provider": "Adzuna" if configured else None,
                "target_role": cached.get("role", role),
                "location": cached.get("location", location),
                "based_on": cached.get("skill_terms", []),
                "jobs": cached.get("jobs", [])[:10],
                "last_checked_at": serialize_document(cached.get("retrieved_at")),
                "cached": True,
                "message": None if configured else "Job search is not configured. Review Adzuna's usage terms, then set ADZUNA_ENABLED=true, ADZUNA_APP_ID, and ADZUNA_APP_KEY on the backend.",
            }
        if not configured:
            return {
                "configured": False,
                "provider": None,
                "jobs": [],
                "last_checked_at": None,
                "cached": False,
                "message": "Job search is not configured. Review Adzuna's usage terms, then set ADZUNA_ENABLED=true, ADZUNA_APP_ID, and ADZUNA_APP_KEY on the backend to enable real listings.",
            }
        if not role:
            return {"configured": True, "provider": "Adzuna", "jobs": [], "last_checked_at": None, "cached": False, "message": "Set a target role in your profile or job preferences to search."}
        country = ADZUNA_COUNTRY
        if not re.fullmatch(r"[a-z]{2}", country):
            raise HTTPException(status_code=503, detail="ADZUNA_COUNTRY must be a two-letter country code.")
        attempts, verified, _ = await dashboard_data(profile)
        dashboard = build_dashboard(profile, attempts, verified, None)
        skill_terms = unique_skill_values([
            *(item.get("skill", "") for item in verified),
            *profile_resume_skills(profile),
            *profile_declared_skills(profile),
            *(item.get("topic", "") for item in sorted(dashboard["skills_to_improve"], key=lambda topic: topic["percentage"])[:2]),
        ])[:5]
        remote = preferences.get("remote")
        cache_key = hashlib.sha256(json.dumps([country, role.casefold(), location.casefold(), remote, prioritize_entry_level, [item.casefold() for item in skill_terms], today], separators=(",", ":")).encode("utf-8")).hexdigest()
        shared_cache = await db["job_provider_cache"].find_one({"cache_key": cache_key}, {"_id": 0})
        if shared_cache:
            jobs = shared_cache.get("jobs", [])
        else:
            jobs = await asyncio.to_thread(fetch_adzuna_jobs, role, location, country, remote, prioritize_entry_level, skill_terms)
            try:
                await db["job_provider_cache"].update_one(
                    {"cache_key": cache_key},
                    {"$setOnInsert": {"cache_key": cache_key, "day_key": today, "country": country, "role": role, "location": location, "skill_terms": skill_terms, "retrieved_at": datetime.now(timezone.utc), "jobs": jobs}},
                    upsert=True,
                )
            except DuplicateKeyError:
                logger.info("Reusing concurrent Adzuna cache entry")
        now = datetime.now(timezone.utc)
        record = {"user_id": user_id, "day_key": today, "retrieved_at": now, "provider": "Adzuna", "role": role, "location": location, "skill_terms": skill_terms, "jobs": jobs}
        try:
            await db["job_recommendations"].update_one(
                {"user_id": user_id, "day_key": today},
                {"$setOnInsert": record},
                upsert=True,
            )
        except DuplicateKeyError:
            logger.info("Reusing concurrent daily job recommendation record")
        stored = await db["job_recommendations"].find_one({"user_id": user_id, "day_key": today}, {"_id": 0}) or record
        return {"configured": True, "provider": "Adzuna", "target_role": role, "location": location, "based_on": skill_terms, "jobs": stored.get("jobs", [])[:10], "last_checked_at": serialize_document(stored.get("retrieved_at")), "cached": False, "message": None}
    except HTTPException:
        raise
    except Exception as exc:
        logger.exception("Job recommendation retrieval failed for student account")
        raise HTTPException(status_code=503, detail="Job recommendations could not be loaded. Please retry later.") from exc


@app.get("/api/jobs/preferences")
async def get_job_preferences(request: Request):
    profile = await require_profile(request)
    if mongo_client is None or mongo_status != "connected":
        raise HTTPException(status_code=503, detail="Job preferences require the connected student database.")
    stored = await mongo_client[MONGODB_DATABASE]["job_preferences"].find_one({"user_id": profile["user_id"]}, {"_id": 0})
    return {"preferences": serialize_document(stored) if stored else {
        "target_role": profile.get("career_goals", {}).get("target_role") or profile.get("career_information", {}).get("target_job_role"),
        "location": profile.get("career_information", {}).get("preferred_work_location"),
        "remote": None,
    }}


@app.put("/api/jobs/preferences")
async def save_job_preferences(request: Request, preferences: JobPreferencesRequest):
    profile = await require_profile(request)
    if mongo_client is None or mongo_status != "connected":
        raise HTTPException(status_code=503, detail="Job preferences require the connected student database.")
    data = preferences.model_dump()
    data["user_id"] = profile["user_id"]
    data["updated_at"] = datetime.now(timezone.utc)
    try:
        await mongo_client[MONGODB_DATABASE]["job_preferences"].update_one(
            {"user_id": profile["user_id"]}, {"$set": data}, upsert=True
        )
        return {"preferences": serialize_document(data)}
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Job preferences could not be saved.") from exc


@app.get("/api/jobs/recommendations")
async def get_job_recommendations(request: Request):
    return await daily_job_recommendations(await require_profile(request))


@app.post("/api/jobs/refresh")
async def refresh_job_recommendations(request: Request):
    return await daily_job_recommendations(await require_profile(request))


@app.get("/api/tests/skill-map")
async def get_skill_map(http_request: Request, target_role: Optional[str] = None):
    profile = await require_profile(http_request)
    return await skill_path_data(profile, target_role)


@app.put("/api/tests/skill-map")
async def select_skill_map(request: SkillMapSelectionRequest, http_request: Request):
    profile = await require_profile(http_request)
    role_name = next(
        (name for name in SKILL_MAP_ROLES if name.casefold() == request.target_role.strip().casefold()),
        None,
    )
    if role_name is None:
        raise HTTPException(status_code=422, detail="Choose one of the available target roles.")
    if mongo_client is None or mongo_status != "connected":
        raise HTTPException(status_code=503, detail="The target role could not be saved because MongoDB is unavailable.")
    try:
        updated = await profiles_collection().update_one(
            {"user_id": profile["user_id"]},
            {"$set": {"career_goals.target_role": role_name, "updated_at": datetime.now(timezone.utc)}},
        )
        if updated.matched_count != 1:
            raise HTTPException(status_code=404, detail="Student profile not found.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=503, detail="The target role could not be saved.") from exc
    profile.setdefault("career_goals", {})["target_role"] = role_name
    return await skill_path_data(profile, role_name)


@app.put("/api/tests/topic-progress")
async def update_topic_progress(request: LearningTopicTaskRequest, http_request: Request):
    profile = await require_profile(http_request)
    skill_key = normalized_skill(request.skill)
    path_key = SKILL_PATH_ALIASES.get(skill_key, skill_key)
    if path_key not in SKILL_PATHS:
        raise HTTPException(status_code=404, detail="Skill learning path not found.")
    topic_definition = next(
        (
            (topic_name, task_names)
            for topic_name, task_names in SKILL_PATHS[path_key]
            if skill_topic_key(topic_name) == skill_topic_key(request.topic)
        ),
        None,
    )
    if topic_definition is None:
        raise HTTPException(status_code=404, detail="Topic learning path not found.")
    topic_name, task_names = topic_definition
    expected_task_ids = {f"task-{index + 1}" for index in range(len(task_names))}
    if request.task_id not in expected_task_ids:
        raise HTTPException(status_code=404, detail="Learning task not found.")
    records = await load_topic_progress(profile["user_id"])
    record = next(
        (
            item for item in records
                if item.get("skill_key") == skill_key
            and item.get("topic_key") == skill_topic_key(topic_name)
        ),
        None,
    )
    existing_tasks = {item.get("task_id"): item for item in (record or {}).get("tasks", [])}
    tasks = [
        {
            "task_id": task_id,
            "completed": request.completed if task_id == request.task_id else bool(existing_tasks.get(task_id, {}).get("completed")),
        }
        for task_id in sorted(expected_task_ids, key=lambda item: int(item.split("-")[1]))
    ]
    progress = {
        **(record or {}),
        "user_id": profile["user_id"],
        "skill": request.skill.strip(),
        "skill_key": skill_key,
        "topic": topic_name,
        "topic_key": skill_topic_key(topic_name),
        "tasks": tasks,
        "updated_at": datetime.now(timezone.utc),
    }
    attempts, _, _ = await dashboard_data(profile)
    assessment = assessment_topic_result(attempts, topic_name)
    progress["assessment_status"] = (
        "not_assessed" if not assessment else "passed" if assessment["passed"] else "needs_improvement"
    )
    progress["assessment_score"] = assessment["percentage"] if assessment else None
    progress["assessment_id"] = assessment["assessment_id"] if assessment else None
    all_tasks_complete = bool(tasks) and all(task["completed"] for task in tasks)
    progress["status"] = (
        "completed" if all_tasks_complete and assessment and assessment["passed"]
        else "in_progress" if request.completed or any(task["completed"] for task in tasks) or assessment
        else "not_started"
    )
    progress["completed_at"] = (
        (record or {}).get("completed_at") or datetime.now(timezone.utc)
        if progress["status"] == "completed"
        else None
    )
    await save_topic_progress_record(profile["user_id"], progress)
    return await skill_path_data(profile)


@app.get("/api/tests/learning-plan")
async def current_learning_plan(http_request: Request):
    profile = await require_profile(http_request)
    user_id = profile["user_id"]
    collection = test_collection()
    if collection is None:
        plans = fallback_learning_plans.get(user_id, [])
        return {"learning_plan": serialize_document(plans[-1]) if plans else None}
    try:
        plan = await mongo_client[MONGODB_DATABASE]["learning_plans"].find_one({"user_id": user_id}, {"_id": 0}, sort=[("created_at", -1)])
        return {"learning_plan": serialize_document(plan) if plan else None}
    except Exception as exc:
        raise HTTPException(status_code=503, detail="The learning plan is temporarily unavailable. Please try again.") from exc


@app.put("/api/tests/learning-plan")
async def update_learning_plan(progress: LearningPlanProgressRequest, http_request: Request):
    profile = await require_profile(http_request)
    user_id = profile["user_id"]
    collection = test_collection()
    if collection is None:
        plans = fallback_learning_plans.get(user_id, [])
        if not plans or progress.task_index >= len(plans[-1].get("tasks", [])):
            raise HTTPException(status_code=404, detail="Learning plan task not found.")
        plans[-1]["tasks"][progress.task_index]["completed"] = progress.completed
        plans[-1]["updated_at"] = datetime.now(timezone.utc)
        return {"learning_plan": serialize_document(plans[-1])}
    try:
        plan = await mongo_client[MONGODB_DATABASE]["learning_plans"].find_one({"user_id": user_id}, sort=[("created_at", -1)])
        if not plan or progress.task_index >= len(plan.get("tasks", [])):
            raise HTTPException(status_code=404, detail="Learning plan task not found.")
        updated_at = datetime.now(timezone.utc)
        await mongo_client[MONGODB_DATABASE]["learning_plans"].update_one(
            {"_id": plan["_id"], "user_id": user_id},
            {"$set": {f"tasks.{progress.task_index}.completed": progress.completed, "updated_at": updated_at}},
        )
        plan["tasks"][progress.task_index]["completed"] = progress.completed
        plan["updated_at"] = updated_at
        return {"learning_plan": serialize_document(plan)}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=503, detail="The learning plan could not be updated. Please try again.") from exc


def calendar_task_id(source: str, identity: str) -> str:
    digest = hashlib.sha256(f"{source}:{identity}".encode("utf-8")).hexdigest()[:24]
    return f"{source}-{digest}"


@app.get("/api/learning/calendar")
async def get_learning_calendar(request: Request):
    profile = await require_profile(request)
    if mongo_client is None or mongo_status != "connected":
        raise HTTPException(status_code=503, detail="The learning calendar requires the connected student database.")
    db = mongo_client[MONGODB_DATABASE]
    saved = await db["learning_calendars"].find_one({"user_id": profile["user_id"]}, {"_id": 0})
    if not saved:
        return {"configured": False, "tasks": [], "preferences": None}
    skill_map = await skill_path_data(profile)
    plan = await db["learning_plans"].find_one({"user_id": profile["user_id"]}, {"_id": 0}, sort=[("created_at", -1)])
    topic_status = {}
    for skill in skill_map.get("skills", []):
        for topic in skill.get("topics", []):
            for task in topic.get("tasks", []):
                key = calendar_task_id("topic", f"{skill['skill']}:{topic['topic']}:{task['task_id']}")
                topic_status[key] = bool(task.get("completed"))
    tasks = []
    today = datetime.now(timezone.utc).date().isoformat()
    for item in saved.get("tasks", []):
        completed = False
        stale = False
        if item.get("source") == "topic":
            completed = topic_status.get(item["task_id"], False)
        elif item.get("source") == "plan":
            if not plan or plan.get("plan_id") != item.get("plan_id"):
                stale = True
            else:
                index = item.get("task_index", -1)
                completed = 0 <= index < len(plan.get("tasks", [])) and bool(plan["tasks"][index].get("completed"))
        status = "stale" if stale else "completed" if completed else "overdue" if item.get("scheduled_date", "") < today else "pending"
        tasks.append({**serialize_document(item), "completed": completed, "stale": stale, "status": status})
    tasks.sort(key=lambda item: (item.get("scheduled_date", ""), item.get("title", "")))
    return {
        "configured": True,
        "preferences": {key: serialize_document(saved.get(key)) for key in ("weekly_minutes", "start_date", "created_at", "updated_at")},
        "tasks": tasks,
    }


@app.put("/api/learning/calendar")
async def create_learning_calendar(request: Request, settings: LearningCalendarRequest):
    profile = await require_profile(request)
    if mongo_client is None or mongo_status != "connected":
        raise HTTPException(status_code=503, detail="The learning calendar requires the connected student database.")
    try:
        start_date = datetime.strptime(settings.start_date, "%Y-%m-%d").date()
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="Choose a valid calendar start date.") from exc
    today = datetime.now(timezone.utc).date()
    if not today <= start_date <= today + timedelta(days=365):
        raise HTTPException(status_code=422, detail="Choose a start date between today and one year from today.")

    skill_map = await skill_path_data(profile)
    db = mongo_client[MONGODB_DATABASE]
    plan = await db["learning_plans"].find_one({"user_id": profile["user_id"]}, {"_id": 0}, sort=[("created_at", -1)])
    candidates = []
    for index, task in enumerate((plan or {}).get("tasks", [])):
        if not task.get("completed"):
            candidates.append({
                "task_id": calendar_task_id("plan", f"{plan['plan_id']}:{index}"),
                "source": "plan",
                "plan_id": plan["plan_id"],
                "task_index": index,
                "title": task.get("task") or f"Review {task.get('topic', 'assessment topic')}",
                "priority": 0 if plan.get("passed") is False else 1,
            })
    for skill in skill_map.get("skills", []):
        for topic in skill.get("topics", []):
            if topic.get("status") == "completed":
                continue
            weak = topic.get("assessment_status") == "needs_improvement"
            for task in topic.get("tasks", []):
                if task.get("completed"):
                    continue
                identity = f"{skill['skill']}:{topic['topic']}:{task['task_id']}"
                candidates.append({
                    "task_id": calendar_task_id("topic", identity),
                    "source": "topic",
                    "skill": skill["skill"],
                    "topic": topic["topic"],
                    "topic_task_id": task["task_id"],
                    "title": task["description"],
                    "priority": 1 if weak else 2,
                })
    candidates.sort(key=lambda item: (item["priority"], item["title"].casefold()))
    tasks_per_week = max(1, math.ceil(settings.weekly_minutes / 30))
    scheduled_tasks = []
    for index, item in enumerate(candidates):
        week, slot = divmod(index, tasks_per_week)
        day_offset = week * 7 + math.floor(slot * 7 / tasks_per_week)
        scheduled_tasks.append({
            **{key: value for key, value in item.items() if key != "priority"},
            "scheduled_date": (start_date + timedelta(days=day_offset)).isoformat(),
        })
    now = datetime.now(timezone.utc)
    document = {
        "user_id": profile["user_id"],
        "weekly_minutes": settings.weekly_minutes,
        "start_date": start_date.isoformat(),
        "tasks": scheduled_tasks,
        "updated_at": now,
    }
    await db["learning_calendars"].update_one(
        {"user_id": profile["user_id"]},
        {"$set": document, "$setOnInsert": {"created_at": now}},
        upsert=True,
    )
    return await get_learning_calendar(request)


@app.put("/api/learning/calendar/task")
async def reschedule_learning_calendar_task(request: Request, update: LearningCalendarRescheduleRequest):
    profile = await require_profile(request)
    if mongo_client is None or mongo_status != "connected":
        raise HTTPException(status_code=503, detail="The learning calendar requires the connected student database.")
    try:
        scheduled_date = datetime.strptime(update.scheduled_date, "%Y-%m-%d").date()
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="Choose a valid task date.") from exc
    today = datetime.now(timezone.utc).date()
    if not today <= scheduled_date <= today + timedelta(days=365):
        raise HTTPException(status_code=422, detail="Reschedule between today and one year from today.")
    collection = mongo_client[MONGODB_DATABASE]["learning_calendars"]
    saved = await collection.find_one({"user_id": profile["user_id"]}, {"_id": 0})
    if not saved or not any(item.get("task_id") == update.task_id for item in saved.get("tasks", [])):
        raise HTTPException(status_code=404, detail="Scheduled learning task not found.")
    tasks = saved["tasks"]
    for item in tasks:
        if item.get("task_id") == update.task_id:
            item["scheduled_date"] = scheduled_date.isoformat()
            break
    await collection.update_one(
        {"user_id": profile["user_id"]},
        {"$set": {"tasks": tasks, "updated_at": datetime.now(timezone.utc)}},
    )
    return await get_learning_calendar(request)


@app.post("/api/auth/login")
async def login_student(credentials: LoginRequest, request: Request, response: Response):
    if mongo_status != "connected":
        raise HTTPException(status_code=503, detail="Sign-in is unavailable while the profile database is disconnected.")
    try:
        profile = await profiles_collection().find_one(
            {"personal_information.email": str(credentials.email).casefold()},
            {"_id": 0},
        )
        if not profile or not verify_password(credentials.password, profile.get("password_hash")):
            raise HTTPException(status_code=401, detail="Email or password is incorrect.")
        token = secrets.token_urlsafe(32)
        now = datetime.now(timezone.utc)
        await profiles_collection().update_one(
            {"user_id": profile["user_id"]},
            {"$set": {
                "auth_token_hash": hashlib.sha256(token.encode("utf-8")).hexdigest(),
                "auth_token_expires_at": now + timedelta(seconds=AUTH_SESSION_TTL_SECONDS),
                "last_login_at": now,
            }},
        )
        set_session_cookie(response, token)
        return {"user_id": profile["user_id"], "profile": serialize_document(profile)}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Sign-in is temporarily unavailable. Please try again.") from exc


@app.get("/api/auth/session")
async def restore_student_session(request: Request, response: Response):
    profile = await require_profile(request)
    if not request.cookies.get("vaishu_session"):
        authorization = request.headers.get("authorization", "")
        if authorization.lower().startswith("bearer "):
            set_session_cookie(response, authorization[7:].strip())
    return {
        "user_id": profile["user_id"],
        "profile": serialize_document(profile),
        "password_login_enabled": bool(profile.get("password_hash")),
    }


@app.post("/api/auth/logout")
async def logout_student(request: Request, response: Response):
    profile = await profile_for_request(request)
    if profile and mongo_status == "connected":
        try:
            await profiles_collection().update_one(
                {"user_id": profile["user_id"]},
                {"$unset": {"auth_token_hash": "", "auth_token_expires_at": ""}},
            )
        except Exception as exc:
            raise HTTPException(status_code=503, detail="Sign-out could not revoke your session. Please try again.") from exc
    clear_session_cookie(response)
    return {"logged_out": True}


@app.post("/api/users/profile", status_code=201)
async def create_profile(
    request: Request,
    response: Response,
    profile_data: str = Form(...),
    password: Optional[str] = Form(default=None, min_length=12, max_length=128),
    resume: Optional[UploadFile] = File(None),
):
    try:
        profile = Profile.model_validate_json(profile_data)
    except (ValidationError, ValueError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=422, detail="Profile data is invalid.") from exc

    now = datetime.now(timezone.utc)
    session_token = secrets.token_urlsafe(32)
    auth_fields = {
        "auth_token_hash": hashlib.sha256(session_token.encode("utf-8")).hexdigest(),
        "auth_token_expires_at": now + timedelta(seconds=AUTH_SESSION_TTL_SECONDS),
    }
    stored_filename = None
    resume_info = {"uploaded": False, "version": None, "original_filename": None, "content_type": None, "size_bytes": 0, "uploaded_at": None, "extraction_status": "not_uploaded", "extracted_information": None}
    resume_version = None
    resume_unchanged = False
    extraction_error: Optional[HTTPException] = None
    try:
        collection = profiles_collection()
        version_collection = mongo_client[MONGODB_DATABASE]["resume_versions"]
        existing = await collection.find_one({"personal_information.email": str(profile.personal_information.email).casefold()})
        authenticated_profile = await profile_for_request(request) if existing else None
        if existing and (not authenticated_profile or authenticated_profile.get("user_id") != existing.get("user_id")):
            raise HTTPException(status_code=409, detail="An account already exists for this email. Sign in to update it.")
        if (not existing or not existing.get("password_hash")) and (not password or len(password) < 12):
            raise HTTPException(status_code=422, detail="Create a password with at least 12 characters to protect this account.")
        user_id = existing["user_id"] if existing else f"usr_{uuid.uuid4().hex[:12]}"

        if resume:
            filename = resume.filename or "resume"
            content = await resume.read()
            extension = validate_resume_content(filename, content)
            file_sha256 = hashlib.sha256(content).hexdigest()
            previous_info = (existing or {}).get("resume_information") or {}
            if previous_info.get("file_sha256") == file_sha256 and previous_info.get("version"):
                resume_info = previous_info
                resume_unchanged = True
            else:
                stored_filename = f"{uuid.uuid4()}{extension}"
                (UPLOAD_DIR / stored_filename).write_bytes(content)
                try:
                    extracted = parse_resume_text(extract_resume_text(extension, content))
                    extraction_status = "extracted"
                except HTTPException as exc:
                    extracted = {section: [] for section in RESUME_SECTIONS}
                    extraction_status = "failed"
                    extraction_error = exc
                version = await next_resume_version(version_collection, user_id)
                resume_version = {"user_id": user_id, "version": version, "uploaded": True, "original_filename": filename, "stored_filename": stored_filename, "content_type": resume.content_type, "size_bytes": len(content), "file_sha256": file_sha256, "uploaded_at": now, "extraction_status": extraction_status, "extracted_information": extracted}
                await version_collection.insert_one(resume_version)
                resume_info = {key: value for key, value in resume_version.items() if key not in {"user_id", "stored_filename"}}

        if existing:
            updates = {**profile.model_dump(), "updated_at": now, **auth_fields}
            if not existing.get("password_hash"):
                if password is None:
                    raise HTTPException(status_code=422, detail="Create a password with at least 12 characters to protect this account.")
                updates["password_hash"] = hash_password(password)
            if resume and not resume_unchanged:
                updates["resume_information"] = resume_info
            await collection.update_one({"user_id": user_id}, {"$set": updates})
            result = {"message": "Profile loaded and updated successfully", "user_id": user_id, "created": False, "resume": public_resume_information(resume_info)}
        else:
            document = {"user_id": user_id, **profile.model_dump(), "password_hash": hash_password(password), "resume_information": resume_info, **auth_fields, "created_at": now, "updated_at": now}
            await collection.insert_one(document)
            result = {"message": "Profile created successfully", "user_id": user_id, "created": True, "resume": public_resume_information(resume_info)}
        set_session_cookie(response, session_token)
        if extraction_error:
            extraction_error.detail = f"{extraction_error.detail} The file was saved as resume version {resume_info['version']}; your previous resume was preserved."
            raise extraction_error
        return result
    except DuplicateKeyError as exc:
        if stored_filename:
            (UPLOAD_DIR / stored_filename).unlink(missing_ok=True)
        raise HTTPException(status_code=409, detail="A profile with this email already exists.") from exc
    except Exception as exc:
        if stored_filename and not extraction_error:
            (UPLOAD_DIR / stored_filename).unlink(missing_ok=True)
        if isinstance(exc, HTTPException):
            raise
        raise HTTPException(status_code=503, detail="Database unavailable. Start MongoDB and try again.") from exc


@app.get("/api/users/{user_id}/resume")
async def get_resume(user_id: str, http_request: Request):
    profile = await require_profile(http_request)
    if profile["user_id"] != user_id:
        raise HTTPException(status_code=403, detail="You can only access your own resume.")
    return {"user_id": user_id, "resume": public_resume_information(profile.get("resume_information") or {})}


@app.get("/api/users/{user_id}/resume/versions")
async def list_resume_versions(user_id: str, http_request: Request):
    profile = await require_profile(http_request)
    if profile["user_id"] != user_id:
        raise HTTPException(status_code=403, detail="You can only access your own resume.")
    try:
        cursor = mongo_client[MONGODB_DATABASE]["resume_versions"].find(
            {"user_id": user_id},
            {"_id": 0, "stored_filename": 0, "user_id": 0},
        ).sort("version", -1)
        versions = [serialize_document(version) async for version in cursor]
        return {"user_id": user_id, "versions": versions}
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Resume history is temporarily unavailable. Please try again.") from exc


@app.put("/api/users/{user_id}/resume")
async def review_resume(user_id: str, review: ResumeReviewRequest, http_request: Request):
    profile = await require_profile(http_request)
    if profile["user_id"] != user_id:
        raise HTTPException(status_code=403, detail="You can only edit your own resume.")
    current = profile.get("resume_information") or {}
    if not current.get("version"):
        raise HTTPException(status_code=404, detail="No uploaded resume was found.")
    if mongo_client is None or mongo_status != "connected":
        raise HTTPException(status_code=503, detail="Resume edits require MongoDB to preserve version history.")
    reviewed_at = datetime.now(timezone.utc)
    extracted = review.extracted_information.model_dump()
    try:
        version_collection = mongo_client[MONGODB_DATABASE]["resume_versions"]
        version = await next_resume_version(version_collection, user_id)
        version_record = {
            key: value for key, value in current.items()
            if key not in {"version", "updated_at", "reviewed_at", "_id", "user_id"}
        }
        version_record.update({
            "user_id": user_id,
            "version": version,
            "prior_version": current["version"],
            "version_type": "student_review",
            "updated_at": reviewed_at,
            "extracted_information": extracted,
            "extraction_status": "reviewed",
            "reviewed_at": reviewed_at,
        })
        await version_collection.insert_one(version_record)
        updated_resume = {
            **current,
            "version": version,
            "prior_version": current["version"],
            "version_type": "student_review",
            "updated_at": reviewed_at,
            "extracted_information": extracted,
            "extraction_status": "reviewed",
            "reviewed_at": reviewed_at,
        }
        updated = await profiles_collection().update_one(
            {"user_id": user_id},
            {"$set": {"resume_information": updated_resume, "updated_at": reviewed_at}},
        )
        if updated.matched_count != 1:
            raise HTTPException(status_code=404, detail="Student profile not found.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=503, detail="The resume review could not be saved. Please try again.") from exc
    return {"message": "Resume details saved as a new version.", "resume": public_resume_information(updated_resume)}


@app.post("/api/users/{user_id}/resume/qualified-skill")
async def add_qualified_skill_to_resume(
    user_id: str,
    request: ResumeSkillUpdateRequest,
    http_request: Request,
    skill: str,
):
    profile = await require_profile(http_request)
    if profile["user_id"] != user_id:
        raise HTTPException(status_code=403, detail="You can only edit your own resume.")
    if not request.confirmed:
        raise HTTPException(status_code=400, detail="Confirm the resume update after reviewing it.")
    if mongo_client is None or mongo_status != "connected":
        raise HTTPException(status_code=503, detail="Resume updates require MongoDB to preserve version history.")
    path = await skill_path_data(profile)
    qualification = next(
        (
            item for item in path["qualified_skills"]
            if normalized_skill(item["skill"]) == normalized_skill(skill)
        ),
        None,
    )
    if qualification is None:
        raise HTTPException(
            status_code=409,
            detail="This skill is not eligible yet. Complete all required topics and pass the skill assessment first.",
        )
    current = profile.get("resume_information") or {}
    current_skills = profile_resume_skills(profile)
    if normalized_skill(skill) in {normalized_skill(item) for item in current_skills}:
        return {
            "message": "This skill is already present in the resume.",
            "added": False,
            "resume": public_resume_information(current),
        }
    extracted = (current.get("extracted_information") or {}).copy()
    extracted.setdefault("education", [])
    extracted["skills"] = unique_skill_values([*(extracted.get("skills") or []), skill])
    for field in ("projects", "certifications", "internships", "work_experience", "achievements"):
        extracted.setdefault(field, [])
    update_date = datetime.now(timezone.utc)
    evidence = {
        "skill": skill,
        "assessment_score": qualification["score"],
        "completion_date": qualification["completed_at"],
        "assessment_id": qualification["assessment_id"],
        "source": "completed_learning_path_and_passed_assessment",
    }
    evidence_records = list(current.get("skill_evidence") or [])
    evidence_records.append(evidence)
    try:
        version_collection = mongo_client[MONGODB_DATABASE]["resume_versions"]
        version = await next_resume_version(version_collection, user_id)
        version_record = {
            key: value for key, value in current.items()
            if key not in {"version", "updated_at", "reviewed_at", "_id", "user_id"}
        }
        version_record.update({
            "user_id": user_id,
            "version": version,
            "uploaded": bool(current.get("uploaded")),
            "version_type": "learning_path_skill_update",
            "prior_version": current.get("version"),
            "updated_at": update_date,
            "source_assessment_id": qualification["assessment_id"],
            "skill_evidence": evidence_records,
            "extracted_information": extracted,
            "extraction_status": "reviewed",
        })
        await version_collection.insert_one(version_record)
        updated_resume = {
            **current,
            "version": version,
            "updated_at": update_date,
            "version_type": "learning_path_skill_update",
            "source_assessment_id": qualification["assessment_id"],
            "skill_evidence": evidence_records,
            "extracted_information": extracted,
            "extraction_status": "reviewed",
        }
        updated = await profiles_collection().update_one(
            {"user_id": user_id},
            {"$set": {"resume_information": updated_resume, "updated_at": update_date}},
        )
        if updated.matched_count != 1:
            raise HTTPException(status_code=404, detail="Student profile not found.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=503, detail="The confirmed resume update could not be saved.") from exc
    return {
        "message": f"{skill} was added to the structured resume profile.",
        "added": True,
        "resume": public_resume_information(updated_resume),
        "download_supported": False,
        "capability_note": "Updated structured resume information is saved and versioned. PDF/DOCX resume generation and download are not currently supported.",
    }



@app.delete("/api/users/{user_id}")
async def delete_profile(user_id: str, http_request: Request):
    profile = await require_profile(http_request)
    if profile["user_id"] != user_id:
        raise HTTPException(status_code=403, detail="You can only access your own profile.")
    try:
        result = await profiles_collection().delete_one(
            {"user_id": user_id}
        )
    except Exception as exc:
        raise HTTPException(
            status_code=503,
            detail="Database unavailable. Start MongoDB and try again."
        ) from exc

    if result.deleted_count == 0:
        raise HTTPException(
            status_code=404,
            detail="User profile not found"
        )

    return {
        "status": "success",
        "message": "User profile deleted successfully",
        "user_id": user_id
    }




@app.get("/api/users")
async def list_profiles(http_request: Request):
    profile = await require_profile(http_request)
    current_user_id = profile["user_id"]
    try:
        cursor = profiles_collection().find(
            {"user_id": current_user_id},
            {
                "_id": 0,
                "user_id": 1,
                "personal_information.full_name": 1,
                "personal_information.email": 1,
                "education.highest_education_level": 1,
                "education.degree_course": 1,
                "career_information.target_job_role": 1,
                "created_at": 1,
            },
        ).sort("created_at", -1)
        profiles = [serialize_document(profile) async for profile in cursor]
        return {"users": profiles, "count": len(profiles)}
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Database unavailable. Start MongoDB and try again.") from exc


@app.get("/api/users/{user_id}")
async def get_profile(user_id: str, http_request: Request):
    profile = await require_profile(http_request)
    if profile["user_id"] != user_id:
        raise HTTPException(status_code=403, detail="You can only access your own profile.")
    try:
        profile = await profiles_collection().find_one({"user_id": user_id})
    except Exception as exc:
        raise HTTPException(status_code=503, detail="Database unavailable. Start MongoDB and try again.") from exc
    if profile is None:
        raise HTTPException(status_code=404, detail="User profile not found.")
    return serialize_document(profile)


app.mount("/", StaticFiles(directory=str(FRONTEND_DIR), html=True), name="frontend")
